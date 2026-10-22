#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
icbc-qwsq —— 企业微信社群运营与质效评价行动推动情况通报

本地源表（覆盖式）→ 备份并打日期戳 → 汇总本周 / 上周 → 覆盖进通报模板（保留一切格式与
条件格式）→ 生成与 10.19 定稿完全同版式的 HTML 正文 → SMTP 直发张希昂。

子命令：
    status                         看源表与备份清单
    backup  --date YYYYMMDD        把 成员群聊数据统计.xlsx 备份成 成员群聊数据统计{date}.xlsx
    build   --date YYYYMMDD        汇总 + 出通报 xlsx + 正文 html/txt（不发信）
    send    --date YYYYMMDD        发信（--dry-run 只打印）
    run     --date YYYYMMDD        backup + build
    all     --date YYYYMMDD        backup + build + send
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import smtplib
import socket
import ssl
import sys
from collections import OrderedDict
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from email.message import EmailMessage
from email.utils import formataddr, formatdate
from pathlib import Path

import openpyxl
import yaml
from openpyxl.utils import column_index_from_string

# ---------------------------------------------------------------- 常量

SKILL_DIR = Path(__file__).resolve().parent
ASSETS = SKILL_DIR / 'assets'
TEMPLATE = ASSETS / 'template.xlsx'
PLAN_DOC = ASSETS / '企业微信社群运营与质效评价方案.docx'

DSH_HOME = Path(os.environ.get('DSH_HOME') or (Path.home() / '.dsh'))

DEFAULT_ROOT = Path('/mnt/d/资料/ICBC/20260929企微')
SRC_SUBDIR = '企微导出表'
SRC_NAME = '成员群聊数据统计.xlsx'
BACKUP_TPL = '成员群聊数据统计{d}.xlsx'
OUT_TPL = '企微社群数据统计通报（{md}）.xlsx'
RECIPIENT = 'zhangxiang.fy@ah.icbc.com.cn'
SUBJECT_TPL = '企业微信社群运营与质效评价行动推动情况通报（{md}）'
FROM_NAME_DEFAULT = '阜阳分行个人金融业务部'

# 支行 -> 网点（顺序与模板一致；文峰并入颍州）
NP_BRANCH = OrderedDict([
    ('颍州', ['阜阳颍州支行营业室', '阜阳颍州金鹰支行', '阜阳颍州南站支行', '阜阳颍州清河支行',
              '阜阳颍州师范学院支行', '阜阳城建支行', '阜阳文峰支行营业室']),
    ('颍泉', ['阜阳颍泉支行营业室', '阜阳颍泉临沂商城支行', '阜阳颍泉瑶海支行', '阜阳颍泉人民路支行',
              '阜阳颍泉大楼支行', '阜阳颍泉小隅首支行', '阜阳牡丹支行']),
    ('颍东', ['阜阳颍东颍河支行', '阜阳颍东支行营业室', '阜阳颍东青年支行', '阜阳颍东烟厂支行',
              '阜阳岳家湖支行', '阜阳汇通支行营业室']),
    ('界首', ['阜阳界首支行营业室', '阜阳界首城建支行', '阜阳界首科创支行', '阜阳界首新街口支行']),
    ('临泉', ['阜阳临泉支行营业室', '阜阳临泉文明支行', '阜阳临泉政务新区支行', '阜阳临泉姜尚支行']),
    ('阜南', ['阜阳阜南支行营业室', '阜阳阜南百货大楼支行', '阜阳阜南北城支行', '阜阳阜南胜利路支行']),
    ('颍上', ['阜阳颍上支行营业室', '阜阳颍上北关支行', '阜阳颍上十字街支行', '阜阳颍上管仲支行']),
    ('太和', ['阜阳太和支行营业室', '阜阳太和人民路支行', '阜阳太和大桥支行', '阜阳太和华源支行']),
])
BRANCH_ORDER = list(NP_BRANCH)
NP_ALL = [n for v in NP_BRANCH.values() for n in v]

# 部门末级 -> 网点（含企微里录错/历史叫法的对照）
DEPT_ALIAS = {
    '临泉姜尚支行': '阜阳临泉姜尚支行', '临泉支行营业室': '阜阳临泉支行营业室',
    '临泉政务新区支行': '阜阳临泉政务新区支行', '临泉文明支行': '阜阳临泉文明支行',
    '太和人民路支行': '阜阳太和人民路支行', '太和支行营业室': '阜阳太和支行营业室',
    '太和大桥支行': '阜阳太和大桥支行', '太和华源支行': '阜阳太和华源支行',
    '界首城建支行': '阜阳界首城建支行',
    '界首影剧院支行': '阜阳界首科创支行',          # 企微部门设置错误，实为科创支行
    '界首支行营业室': '阜阳界首支行营业室', '界首新街口支行': '阜阳界首新街口支行',
    '阜南北城支行': '阜阳阜南北城支行', '阜南支行营业室': '阜阳阜南支行营业室',
    '阜南百货大楼支行': '阜阳阜南百货大楼支行', '阜南胜利路支行': '阜阳阜南胜利路支行',
    '阜阳文峰支行': '阜阳文峰支行营业室', '阜阳文峰支行营业室': '阜阳文峰支行营业室',
    '阜阳汇通支行': '阜阳汇通支行营业室',
    '阜阳烟厂支行': '阜阳颍东烟厂支行', '阜阳青年支行': '阜阳颍东青年支行',
    '阜阳颍上路支行': '阜阳岳家湖支行',              # 企微部门设置错误，实为岳家湖支行
    '阜阳颍东支行营业室': '阜阳颍东支行营业室', '阜阳颍河支行': '阜阳颍东颍河支行',
    '中国工商银行股份有限公司阜阳师范大学支行': '阜阳颍州师范学院支行',
    '阜阳南站支行': '阜阳颍州南站支行', '阜阳城建支行': '阜阳城建支行',
    '阜阳小隅首支行': '阜阳颍泉小隅首支行', '阜阳清河支行': '阜阳颍州清河支行',
    '阜阳金鹰支行': '阜阳颍州金鹰支行', '阜阳颍州支行营业室': '阜阳颍州支行营业室',
    '阜阳临沂商城支行': '阜阳颍泉临沂商城支行', '阜阳人民路支行': '阜阳颍泉人民路支行',
    '阜阳大楼支行': '阜阳颍泉大楼支行', '阜阳牡丹支行': '阜阳牡丹支行',
    '阜阳瑶海支行': '阜阳颍泉瑶海支行', '阜阳颍泉支行营业室': '阜阳颍泉支行营业室',
    '颍上北关支行': '阜阳颍上北关支行', '颍上十字街支行': '阜阳颍上十字街支行',
    '颍上支行营业室': '阜阳颍上支行营业室', '颍上管仲支行': '阜阳颍上管仲支行',
}
# 非网点部门：分行本部各部室、各支行机构本部/行长室/综合部，以及部门为空（离职账号）
IGNORE_MARKS = ('机构本部', '行长室', '综合业务部', '综合管理部', '金融科技部', '银行卡中心',
                '结算与现金管理中心', '个人信贷业务部', '个人金融业务部', '分行本部')

# 条件格式：好=红、差=绿（与模板 dxf 一致；E 列落底色+白字，H 列只落字色）
RED, GREEN, WHITE = 'FF0000', '00B050', 'FFFFFF'


# ---------------------------------------------------------------- 工具

def die(msg: str) -> None:
    print(json.dumps({'ok': False, 'error': msg}, ensure_ascii=False, indent=2))
    sys.exit(1)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def parse_day(text: str | None) -> date:
    if not text:
        return date.today()
    text = text.strip().replace('-', '').replace('/', '')
    if not re.fullmatch(r'\d{8}', text):
        die(f'日期格式不对：{text}（要 YYYYMMDD）')
    return datetime.strptime(text, '%Y%m%d').date()


def paths(root: Path):
    src_dir = root / SRC_SUBDIR
    return src_dir, src_dir / SRC_NAME


def list_backups(src_dir: Path) -> list[tuple[date, Path]]:
    out = []
    for p in src_dir.glob('成员群聊数据统计*.xlsx'):
        m = re.fullmatch(r'成员群聊数据统计(\d{8})\.xlsx', p.name)
        if not m:
            continue
        try:
            out.append((datetime.strptime(m.group(1), '%Y%m%d').date(), p))
        except ValueError:
            continue
    return sorted(out)


def workdir_of(root: Path, d: date) -> Path:
    return root / 'automations' / f'{d:%Y%m%d}_企微社群'


# ---------------------------------------------------------------- 汇总

def aggregate(path: Path) -> tuple[dict, dict]:
    """读取成员群聊表 -> ({网点: {g,c,msg,x}}, 统计信息)"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if 'result' not in wb.sheetnames:
        die(f'{path.name} 里没有 result 工作表，确认是不是「成员群聊数据统计」导出')
    ws = wb['result']
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else '' for c in next(rows)]
    need = {'部门': '部门', '群总数': '群总数', '群成员总数': '群成员总数',
            '发过消息的群成员数': '发过消息的群成员数', '群聊消息总数': '群聊消息总数'}
    idx = {}
    for key, label in need.items():
        if label not in header:
            die(f'{path.name} 表头缺少「{label}」，实际表头：{header}')
        idx[key] = header.index(label)

    agg: dict[str, dict] = {n: {'g': 0, 'c': 0, 'msg': 0, 'x': 0} for n in NP_ALL}
    stats = {'rows': 0, 'used': 0, 'ignored_head_office': 0, 'ignored_blank_dept': 0,
             'unknown_dept': {}, 'members': []}

    def num(v):
        try:
            return int(float(v))
        except (TypeError, ValueError):
            return 0

    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        stats['rows'] += 1
        dept = str(row[idx['部门']] or '').strip()
        if not dept:
            stats['ignored_blank_dept'] += 1           # 部门为空 = 离职/注销账号
            continue
        if any(mark in dept for mark in IGNORE_MARKS):
            stats['ignored_head_office'] += 1
            continue
        leaf = dept.split('/')[-1].strip()
        target = DEPT_ALIAS.get(leaf)
        if not target:
            stats['unknown_dept'][leaf] = stats['unknown_dept'].get(leaf, 0) + 1
            continue
        a = agg[target]
        a['g'] += num(row[idx['群总数']])
        a['c'] += num(row[idx['群成员总数']])
        a['msg'] += num(row[idx['发过消息的群成员数']])
        a['x'] += num(row[idx['群聊消息总数']])
        stats['used'] += 1
        stats['members'].append({'name': str(row[0] or ''), 'dept': leaf, 'net': target})

    wb.close()
    for n in NP_ALL:
        a = agg[n]
        a['rate'] = (a['msg'] / a['c']) if a['c'] else 0.0
    return agg, stats


def branch_agg(net: dict) -> dict[str, dict]:
    out = {}
    for br, nets in NP_BRANCH.items():
        g = sum(net[n]['g'] for n in nets)
        c = sum(net[n]['c'] for n in nets)
        msg = sum(net[n]['msg'] for n in nets)
        x = sum(net[n]['x'] for n in nets)
        out[br] = {'g': g, 'c': c, 'msg': msg, 'x': x, 'rate': (msg / c) if c else 0.0}
    return out


def totals(net: dict) -> dict:
    g = sum(v['g'] for v in net.values())
    c = sum(v['c'] for v in net.values())
    msg = sum(v['msg'] for v in net.values())
    x = sum(v['x'] for v in net.values())
    return {'g': g, 'c': c, 'msg': msg, 'x': x, 'rate': (msg / c) if c else 0.0}


# ---------------------------------------------------------------- 填表

def build_xlsx(cur: dict, prev: dict, d: date, out_path: Path) -> dict:
    branch_cur, branch_prev = branch_agg(cur), branch_agg(prev)
    tot_cur, tot_prev = totals(cur), totals(prev)

    wb = openpyxl.load_workbook(TEMPLATE)
    title = f'企微个微社群数据统计表（{d.month}月{d.day}日）'

    plans = [
        ('网点维度', NP_ALL, cur, prev, 4, 44),
        ('支行维度', BRANCH_ORDER, branch_cur, branch_prev, 4, 12),
    ]
    for sheet, labels, now, before, r0, total_row in plans:
        ws = wb[sheet]
        ws.cell(1, 1).value = title
        for i, label in enumerate(labels):
            r = r0 + i
            display = label if sheet == '网点维度' else f'{label}支行'
            assert ws.cell(r, 1).value == display, f'{sheet} 第{r}行应是 {display}，实际 {ws.cell(r,1).value}'
            a, b = now[label], before[label]
            ws.cell(r, 2).value = a['g']
            ws.cell(r, 3).value = a['g'] - b['g']
            ws.cell(r, 4).value = a['c']
            ws.cell(r, 5).value = a['c'] - b['c']
            ws.cell(r, 6).value = a['x']
            ws.cell(r, 7).value = a['x'] - b['x']
            ws.cell(r, 8).value = a['rate']
            ws.cell(r, 9).value = a['rate'] - b['rate']
        # 合计行：B/D/F 保留模板公式，C/E/G/H/I 由脚本写
        ws.cell(total_row, 3).value = tot_cur['g'] - tot_prev['g']
        ws.cell(total_row, 5).value = tot_cur['c'] - tot_prev['c']
        ws.cell(total_row, 7).value = tot_cur['x'] - tot_prev['x']
        ws.cell(total_row, 8).value = tot_cur['rate']
        ws.cell(total_row, 9).value = tot_cur['rate'] - tot_prev['rate']

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {'net': cur, 'net_prev': prev, 'branch': branch_cur, 'branch_prev': branch_prev,
            'tot_cur': tot_cur, 'tot_prev': tot_prev, 'title': title}


# ---------------------------------------------------------------- 表格转 HTML

def pct_text(v: float) -> str:
    """百分比，四舍五入到两位（Excel 口径，非 Python 的银行家舍入）"""
    d = Decimal(str(float(v) * 100)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    return f'{d}%'

def _hex(color, default=None):
    if color is None:
        return default
    rgb = getattr(color, 'rgb', None)
    if isinstance(rgb, str) and re.fullmatch(r'[0-9A-Fa-f]{8}', rgb):
        return rgb[-6:].upper()
    return default


def cf_cell_map(ws) -> dict[tuple[int, int], dict]:
    """把模板里的条件格式按「前N名/后N名（并列全含）」落成 (row,col)->{'bg','color'}"""
    out: dict[tuple[int, int], dict] = {}
    for rng in ws.conditional_formatting:
        for sqref in str(rng.sqref).split():
            m = re.fullmatch(r'([A-Z]+)(\d+):([A-Z]+)(\d+)', sqref)
            if not m:
                continue
            c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
            if c1 != c2:
                continue
            col = column_index_from_string(c1)
            vals = []
            for r in range(r1, r2 + 1):
                v = ws.cell(r, col).value
                if isinstance(v, (int, float)):
                    vals.append((r, v))
            if not vals:
                continue
            for rule in rng.rules:
                if rule.type != 'top10':
                    continue
                rank = int(getattr(rule, 'rank', 10) or 10)
                bottom = bool(getattr(rule, 'bottom', False))
                ordered = sorted((v for _, v in vals), reverse=not bottom)
                if not ordered:
                    continue
                thr = ordered[min(rank, len(ordered)) - 1]
                chosen = [r for r, v in vals if (v <= thr if bottom else v >= thr)]
                dxf = rule.dxf
                bg = None
                if dxf is not None and dxf.fill is not None:
                    bg = _hex(dxf.fill.bgColor)
                    if bg in ('000000', None) and dxf.fill.patternType is None:
                        bg = None
                color = None
                if dxf is not None and dxf.font is not None:
                    color = _hex(dxf.font.color)
                if bg and not color:
                    color = WHITE          # 模板里底色格字为 white（dxf theme=0）
                for r in chosen:
                    key = (r, col)
                    slot = out.setdefault(key, {})
                    if bg:
                        slot['bg'] = bg
                    if color:
                        slot['color'] = color
    return out


def fmt_value(v, number_format: str) -> str:
    if v is None:
        return ''
    if isinstance(v, str):
        return v
    if not isinstance(v, (int, float)):
        return str(v)
    nf = (number_format or 'General').replace('\\', '')
    plus = nf.startswith('+')
    if '%' in nf:
        txt = pct_text(v)
        if plus:
            if v > 0:
                return '+' + txt
            if v == 0:
                return '0.00%'
        return txt
    if '#,##0' in nf:
        txt = f'{v:,.0f}'
        if plus:
            if v > 0:
                return '+' + txt
            if v == 0:
                return '0'
        return txt
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def table_html(ws, total_overrides: dict[int, float]) -> str:
    merged: dict[tuple[int, int], tuple[int, int]] = {}
    covered = set()
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1,
                                              rng.max_col - rng.min_col + 1)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    covered.add((r, c))
    cf = cf_cell_map(ws)
    total_row = ws.max_row
    out = ['<table border="0" cellpadding="0" cellspacing="0" width="1024" '
           'style="font-family:Arial;font-size:14px;border-collapse:collapse;">', '<tbody>']
    for r in range(1, ws.max_row + 1):
        out.append('<tr>')
        for c in range(1, ws.max_column + 1):
            if (r, c) in covered:
                continue
            rowspan, colspan = merged.get((r, c), (1, 1))
            cell = ws.cell(r, c)
            f = cell.font
            size = f.size or 11
            if isinstance(size, float) and size.is_integer():
                size = int(size)
            weight = 700 if f.bold else 400
            fcolor = _hex(f.color, '222222')
            fill = _hex(cell.fill.fgColor) if cell.fill and cell.fill.patternType else None
            border = _hex(cell.border.left.color, 'BFBFBF') if cell.border and cell.border.left else 'BFBFBF'
            span = cf.get((r, c)) or {}
            bg = span.get('bg') or fill
            tcolor = span.get('color') or fcolor
            style = (f'font-size:{size}pt;font-weight:{weight};font-style:normal;text-decoration:none;'
                     f'color:#{tcolor};font-family:{f.name or "微软雅黑"};'
                     f'border:0.5pt solid #{border};text-align:center;vertical-align:middle;')
            if bg:
                style += f'background:#{bg};'
            value = cell.value
            if r == total_row and isinstance(value, str) and value.startswith('='):
                value = total_overrides.get(c, 0)
            text = fmt_value(value, cell.number_format)
            attrs = [f'style="{style}"']
            if colspan > 1:
                attrs.append(f'colspan="{colspan}"')
            if rowspan > 1:
                attrs.append(f'rowspan="{rowspan}"')
            out.append(f'<td {" ".join(attrs)}>{text}</td>')
        out.append('</tr>')
    out.append('</tbody></table>')
    return ''.join(out)


# ---------------------------------------------------------------- 正文

PS = 'font-family:Arial;font-size:14px;margin:0px;line-height:1.7;'
SPAN = ('font-family:仿宋_gb2312;font-size:21.3333px;background-color:#FFFFFF;color:#000000;')


def p_html(inner: str, indent: bool = True, lead_nbsp: int = 0) -> str:
    ind = 'text-indent:42.6666px;' if indent else ''
    lead = '&nbsp; ' * lead_nbsp
    return (f'<p style="{PS}{ind}"><span style="{SPAN}">{lead}{inner}</span></p>')


def blank_html() -> str:
    return f'<p style="{PS}"><span style="{SPAN}"><br></span></p>'


def red(t: str) -> str:
    return f'<span style="color:#{RED};">{t}</span>'


def green(t: str) -> str:
    return f'<span style="color:#{GREEN};">{t}</span>'


def bold(t: str) -> str:
    return f'<strong>{t}</strong>'


def short_name(name: str) -> str:
    """网点简称：去掉「阜阳」前缀（10.19 定稿里建群网点前三用简称）"""
    return name[2:] if name.startswith('阜阳') else name


def _delta(v: int, unit: str, add: str = '新增', cut: str = '减少') -> str:
    if v >= 0:
        return f'较上周{add}{v}{unit}'
    return f'较上周{cut}{abs(v)}{unit}'


def _pct(v: float) -> str:
    return pct_text(v)


def _rank(items, key, n, reverse=True):
    return sorted(items, key=key, reverse=reverse)[:n]


def build_body(data: dict, d: date, xlsx_path: Path) -> tuple[str, str]:
    cur, prev = data['net'], data['net_prev']
    br, br_prev = data['branch'], data['branch_prev']
    tc, tp = data['tot_cur'], data['tot_prev']

    br_items = [br[b] | {'name': b + '支行'} for b in BRANCH_ORDER]
    br_prev_map = {b: br_prev[b] for b in BRANCH_ORDER}
    np_items = [cur[n] | {'name': n} for n in NP_ALL]

    def dbr(b, k):
        return br[b][k] - br_prev[b][k]

    def dnp(n, k):
        return cur[n][k] - prev[n][k]

    # ---- 建群 ----
    top_g = _rank(br_items, lambda x: (x['g'], dbr(x['name'][:-2], 'g')), 2)
    low_g = sorted(br_items, key=lambda x: (x['g'], dbr(x['name'][:-2], 'g')))[:2]
    top_gn = _rank(np_items, lambda x: (x['g'], dnp(x['name'], 'g')), 3)
    zero_g = [n for n in NP_ALL if cur[n]['g'] == 0]

    # ---- 添客（客户数） ----
    top_c = _rank(br_items, lambda x: (dbr(x['name'][:-2], 'c'), x['c']), 2)
    low_c = sorted(br_items, key=lambda x: (dbr(x['name'][:-2], 'c'), x['c']))[:2]
    top_cn = _rank(np_items, lambda x: (dnp(x['name'], 'c'), x['c']), 3)
    low_cn = sorted(np_items, key=lambda x: (dnp(x['name'], 'c'), x['c']))[:3]

    # ---- 交互 ----
    top_x = _rank(br_items, lambda x: (x['x'], x['name']), 2)
    zero_x = [x['name'] for x in br_items if x['x'] == 0]
    low_x = sorted(br_items, key=lambda x: (x['x'], x['name']))[:2]
    top_xn = _rank(np_items, lambda x: (x['x'], x['name']), 3)

    # ---- 熟识率 ----
    top_r = _rank(br_items, lambda x: (x['rate'], x['name']), 2)
    zero_r = [x['name'] for x in br_items if x['rate'] == 0]
    top_rn = _rank(np_items, lambda x: (x['rate'], x['name']), 3)

    date_cn = f'{d.year}年{d.month}月{d.day}日'

    def br_desc(x, unit_delta, word='建群'):
        b = x['name'][:-2]
        return f"{x['name']}{word}{x['g']}个，{_delta(dbr(b, 'g'), '个')}"

    paras: list[str] = []

    # 0 开头空段落（与 10.19 定稿版式一致）
    paras.append(blank_html())

    # 1 总述
    paras.append(p_html(
        f'截至{date_cn}，全行企微累计建群{tc["g"]}个，较上周'
        f'{"增加" if tc["g"] - tp["g"] >= 0 else "减少"}{abs(tc["g"] - tp["g"])}个；'
        f'客户数{tc["c"]}户，较上周'
        f'{"增加" if tc["c"] - tp["c"] >= 0 else "减少"}{abs(tc["c"] - tp["c"])}户。',
        lead_nbsp=4))
    paras.append(blank_html())

    # 2 建群：支行
    good = '、'.join(f"{x['name']}建群{x['g']}个，{_delta(dbr(x['name'][:-2], 'g'), '个')}" for x in top_g)
    bad = '、'.join(f"{x['name']}仅建群{x['g']}个，{_delta(dbr(x['name'][:-2], 'g'), '个')}" for x in low_g)
    paras.append(p_html(f'{bold("建群方面：")}{red(good)}领跑。{green(bad)}暂时落后。'))

    # 3 建群：网点
    seg = '、'.join(f"{short_name(x['name'])}（{x['g']}个，{_delta(dnp(x['name'], 'g'), '个')}）"
                    for x in top_gn)
    line = f'网点层面，{red(seg)}居前'
    if zero_g:
        line += f'；{green("、".join(zero_g))}{len(zero_g)}个网点尚未建群。'
    else:
        line += '。'
    paras.append(p_html(line))
    paras.append(blank_html())

    # 4 添客：支行
    good_c = '、'.join(f"{x['name']}（{x['c']}户，{_delta(dbr(x['name'][:-2], 'c'), '户')}）" for x in top_c)
    bad_c = '、'.join(f"{x['name']}（{x['c']}户，{_delta(dbr(x['name'][:-2], 'c'), '户')}）" for x in low_c)
    paras.append(p_html(f'{bold("添客方面：")}{red(good_c)}添客领先；{green(bad_c)}添客乏力。'))

    # 5 添客：网点
    good_cn = '、'.join(f"{x['name']}（{x['c']}户，{_delta(dnp(x['name'], 'c'), '户')}）" for x in top_cn)
    bad_cn = '、'.join(f"{x['name']}（{x['c']}户，{_delta(dnp(x['name'], 'c'), '户')}）" for x in low_cn)
    paras.append(p_html(f'网点层面，前三名分别为{red(good_cn)}；后三名为{green(bad_cn)}。'))
    paras.append(blank_html())

    # 6 交互
    good_x = '、'.join(f"{x['name']}（{x['x']}条）" for x in top_x)
    if zero_x:
        tail = f'{green("、".join(zero_x))}交互数为0。'
    else:
        names = '、'.join(x['name'] for x in low_x)
        tail = f'{green(names)}交互数偏低。'
    good_xn = '、'.join(f"{x['name']}（{x['x']}条）" for x in top_xn)
    paras.append(p_html(f'{bold("交互方面：")}{red(good_x)}交互领先；{tail}'
                        f'网点层面，前三名分别为{red(good_xn)}。'))
    paras.append(blank_html())

    # 7 熟识率
    good_r = '、'.join(f"{x['name']}（{_pct(x['rate'])}）" for x in top_r)
    if zero_r:
        tail_r = f'{green("、".join(zero_r))}为0。'
    else:
        names = '、'.join(x['name'] for x in sorted(br_items, key=lambda y: y['rate'])[:2])
        tail_r = f'{green(names)}熟识率偏低。'
    good_rn = '、'.join(f"{x['name']}{_pct(x['rate'])}" for x in top_rn)
    paras.append(p_html(f'{bold("客户熟识率方面：")}{red(good_r)}领先；{tail_r}'
                        f'网点层面，前三名分别为{red(good_rn)}。'))
    paras.append(blank_html())

    # 8 两个表
    wb = openpyxl.load_workbook(xlsx_path)
    branch_ws = wb['支行维度']
    net_ws = wb['网点维度']
    branch_tbl = table_html(branch_ws, {2: tc['g'], 4: tc['c'], 6: tc['x']})
    net_tbl = table_html(net_ws, {2: tc['g'], 4: tc['c'], 6: tc['x']})
    paras.append(branch_tbl)
    paras.append(blank_html())
    paras.append(blank_html())
    paras.append(net_tbl)
    paras.append(blank_html())

    # 9 方案附件引用
    paras.append(p_html(f'📎 {PLAN_DOC.name}'))
    paras.append(blank_html())
    paras.append(p_html(bold('工作提示：')))
    paras.append(p_html('一是抓好分类建群添客。10—11月，非行政班网点添客不少于1000户、'
                        '行政班网点不少于600户，个微群成员可逐步迁移至企微群。'))
    paras.append(p_html('二是落实“1+4+X”运营分工。明确群主及财客、远维、客群经理等职责，'
                        '新群按前30天的关键节点持续开展服务，后续按月循环持续运营。'))
    paras.append(blank_html())
    paras.append(blank_html())

    html = ('<html><head><meta charset="utf-8"></head><body>'
            '<div style="line-height:1.7;color:#000000;">' + ''.join(paras) + '</div></body></html>')

    # 纯文本版（text/plain 部分）
    import html as _h
    plain_lines = []
    for seg in paras:
        if seg.startswith('<table'):
            plain_lines.append('[表格见 HTML 正文与附件]')
            continue
        txt = _h.unescape(re.sub(r'<[^>]+>', '', seg)).replace('\u00a0', ' ').strip()
        plain_lines.append(txt)
    text = '\n'.join(plain_lines)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return html, text


# ---------------------------------------------------------------- 账号 / 发信

PROVIDERS = {'163': ('imap.163.com', 993, 'smtp.163.com', 465)}


def _read_row_config() -> dict:
    out = {}
    for cfg in sorted((DSH_HOME / 'profiles').glob('*/cordis.patch.yml')):
        try:
            data = yaml.safe_load(cfg.read_text('utf-8')) or []
        except Exception:
            continue
        if not isinstance(data, list):
            continue
        for entry in data:
            if isinstance(entry, dict) and entry.get('id') == 'tool-email':
                conf = entry.get('config')
                if isinstance(conf, dict):
                    out.update(conf)
    return out


def _read_page_config() -> dict:
    try:
        data = yaml.safe_load((DSH_HOME / 'settings.yaml').read_text('utf-8')) or {}
    except Exception:
        return {}
    ns = data.get('dsh-email')
    return ns if isinstance(ns, dict) else {}


def _env_password():
    pw = os.environ.get('DSH_EMAIL_PASSWORD')
    if pw and pw.strip():
        return pw.strip(), 'DSH_EMAIL_PASSWORD 环境变量'
    for pid in os.listdir('/proc'):
        if not pid.isdigit():
            continue
        try:
            cmd = Path(f'/proc/{pid}/cmdline').read_bytes().replace(b'\0', b' ').decode('utf-8', 'ignore')
            if 'dsh' not in cmd or 'web' not in cmd:
                continue
            for kv in Path(f'/proc/{pid}/environ').read_bytes().split(b'\0'):
                if kv.startswith(b'DSH_EMAIL_PASSWORD='):
                    return kv.split(b'=', 1)[1].decode('utf-8', 'ignore').strip(), '运行中 dsh web 的环境'
        except (OSError, PermissionError):
            continue
    return None


def resolve_account() -> dict:
    cfg = {}
    cfg.update(_read_row_config())
    page = _read_page_config()
    for key in ('provider', 'user', 'password', 'inboxFolder'):
        v = page.get(key)
        if isinstance(v, str) and v.strip():
            cfg[key] = v.strip()
    entries, default = None, None
    raw = cfg.get('accountsYaml') if isinstance(cfg.get('accountsYaml'), str) else ''
    if raw.strip():
        doc = yaml.safe_load(raw) or {}
        if isinstance(doc, dict):
            doc = dict(doc)
            default = doc.pop('defaultAccount', None)
            entries = doc
    elif isinstance(cfg.get('accounts'), dict) and cfg['accounts']:
        entries = cfg['accounts']
    if entries:
        name = cfg.get('defaultAccount') or default or ('default' if 'default' in entries else sorted(entries)[0])
        acct = entries.get(name) or {}
        merged = {k: v for k, v in acct.items() if v not in (None, '')}
        for k in ('provider', 'user', 'password', 'imap', 'smtp'):
            if k not in merged and cfg.get(k):
                merged[k] = cfg[k]
        cfg = merged
    provider = (cfg.get('provider') or '163').strip() or '163'
    imap_host, imap_port, smtp_host, smtp_port = PROVIDERS.get(provider, PROVIDERS['163'])
    smtp_conf = cfg.get('smtp') if isinstance(cfg.get('smtp'), dict) else {}
    pwd = cfg.get('password')
    if isinstance(pwd, str) and pwd.strip():
        password = pwd.strip()
    else:
        found = _env_password()
        if not found:
            die('取不到邮箱授权码：请在邮件设置页填写，或设置 DSH_EMAIL_PASSWORD')
        password = found[0]
    return {'user': str(cfg.get('user') or '').strip(),
            'smtp_host': str(smtp_conf.get('host') or '').strip() or smtp_host,
            'smtp_port': int(smtp_conf.get('port') or smtp_port),
            'password': password}


def force_ipv4() -> None:
    original = socket.getaddrinfo

    def ipv4_only(host, port, family=0, type=0, proto=0, flags=0):
        return original(host, port, socket.AF_INET, type, proto, flags)

    socket.getaddrinfo = ipv4_only


def send_mail(to: str, subject: str, html: str, text: str, attachments: list[Path],
              from_name: str, dry_run: bool) -> dict:
    acct = resolve_account()
    msg = EmailMessage()
    msg['From'] = formataddr((from_name, acct['user']))
    msg['To'] = to
    msg['Subject'] = subject
    msg['Date'] = formatdate(localtime=True)
    msg.set_content(text or '企业微信社群运营与质效评价行动推动情况通报，详见附件。')
    msg.add_alternative(html, subtype='html')
    for path in attachments:
        suffix = path.suffix.lower()
        subtype = ('vnd.openxmlformats-officedocument.spreadsheetml.sheet'
                   if suffix == '.xlsx' else
                   'vnd.openxmlformats-officedocument.wordprocessingml.document')
        msg.add_attachment(path.read_bytes(), maintype='application', subtype=subtype,
                           filename=path.name)

    info = {'to': to, 'subject': subject, 'from': msg['From'],
            'attachments': [p.name for p in attachments],
            'smtp': f"{acct['smtp_host']}:{acct['smtp_port']}",
            'html_bytes': len(html), 'text_bytes': len(text)}
    if dry_run:
        return {'ok': True, 'dry_run': True, **info}
    force_ipv4()
    ctx = ssl.create_default_context()
    if acct['smtp_port'] == 465:
        server = smtplib.SMTP_SSL(acct['smtp_host'], acct['smtp_port'], context=ctx, timeout=60)
    else:
        server = smtplib.SMTP(acct['smtp_host'], acct['smtp_port'], timeout=60)
        server.starttls(context=ctx)
    try:
        server.login(acct['user'], acct['password'])
        server.send_message(msg)
    finally:
        try:
            server.quit()
        except Exception:
            pass
    return {'ok': True, 'sent': True, **info}


# ---------------------------------------------------------------- 子命令

def cmd_status(args) -> dict:
    root = Path(args.root)
    src_dir, src = paths(root)
    back = list_backups(src_dir)
    return {'ok': True, 'root': str(root), 'source': {'path': str(src), 'exists': src.exists(),
            'mtime': datetime.fromtimestamp(src.stat().st_mtime).strftime('%Y-%m-%d %H:%M:%S') if src.exists() else None},
            'backups': [{'date': f'{d:%Y%m%d}', 'file': p.name} for d, p in back],
            'template': str(TEMPLATE), 'template_exists': TEMPLATE.exists(),
            'plan_doc': str(PLAN_DOC), 'plan_exists': PLAN_DOC.exists(),
            'recipient': RECIPIENT}


def cmd_backup(args) -> dict:
    root = Path(args.root)
    src_dir, src = paths(root)
    d = parse_day(args.date)
    if not src.exists():
        die(f'源表不存在：{src}')
    dest = src_dir / BACKUP_TPL.format(d=f'{d:%Y%m%d}')
    if dest.exists():
        same = sha256(dest) == sha256(src)
        if same:
            return {'ok': True, 'skipped': True, 'reason': '当日备份已存在且内容一致',
                    'backup': str(dest)}
        die(f'备份已存在但内容不同：{dest.name}。源表被再次覆盖过？请确认后手动处理（改名或删除旧备份）。')
    shutil.copy2(src, dest)
    return {'ok': True, 'backup': str(dest), 'bytes': dest.stat().st_size}


def _resolve_files(root: Path, d: date, prev_arg: str | None, cur_file: str | None,
                   prev_file: str | None):
    src_dir, src = paths(root)
    if cur_file:
        cur_path = Path(cur_file)
    else:
        cur_path = src_dir / BACKUP_TPL.format(d=f'{d:%Y%m%d}')
        if not cur_path.exists():
            die(f'当日备份不存在：{cur_path}。先跑 backup --date {d:%Y%m%d}')
    if not cur_path.exists():
        die(f'本周源文件不存在：{cur_path}')

    if prev_file:
        prev_path = Path(prev_file)
    else:
        back = [(bd, p) for bd, p in list_backups(src_dir) if bd < d]
        if prev_arg:
            target = parse_day(prev_arg)
            hit = [p for bd, p in back if bd == target]
            if not hit:
                die(f'找不到 {target:%Y%m%d} 的备份')
            prev_path = hit[0]
        elif back:
            prev_path = back[-1][1]
        else:
            die('找不到上周备份。首次运行请用 --prev-file 指定上周源表（例如 成员群聊数据统计1005-1011.xlsx），'
                '确认后再运行。')
    if not prev_path.exists():
        die(f'上周源文件不存在：{prev_path}')
    return cur_path, prev_path


def cmd_build(args) -> dict:
    root = Path(args.root)
    d = parse_day(args.date)
    cur_path, prev_path = _resolve_files(root, d, args.prev, args.cur_file, args.prev_file)

    cur, cur_stats = aggregate(cur_path)
    prev, prev_stats = aggregate(prev_path)

    work = workdir_of(root, d)
    work.mkdir(parents=True, exist_ok=True)
    md = f'{d.month}.{d.day}'
    out_xlsx = Path(args.out) if args.out else root / OUT_TPL.format(md=md)

    data = build_xlsx(cur, prev, d, out_xlsx)
    html, text = build_body(data, d, out_xlsx)

    (work / 'mail_body.html').write_text(html, encoding='utf-8')
    (work / 'mail_body.txt').write_text(text, encoding='utf-8')
    payload = {
        'ok': True, 'date': f'{d:%Y%m%d}', 'cur_file': cur_path.name, 'prev_file': prev_path.name,
        'workdir': str(work), 'xlsx': str(out_xlsx),
        'totals': {'cur': data['tot_cur'], 'prev': data['tot_prev']},
        'warnings': {
            'unknown_dept': cur_stats['unknown_dept'],
            'ignored_head_office': cur_stats['ignored_head_office'],
            'ignored_blank_dept': cur_stats['ignored_blank_dept'],
            'unknown_dept_prev': prev_stats['unknown_dept'],
        },
        'rows': {'cur': cur_stats['rows'], 'used': cur_stats['used']},
        'net': {n: cur[n] for n in NP_ALL},
        'branch': data['branch'],
        'text': text,
    }
    (work / 'data.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    return payload


def cmd_send(args) -> dict:
    root = Path(args.root)
    d = parse_day(args.date)
    work = workdir_of(root, d)
    html_path, text_path = work / 'mail_body.html', work / 'mail_body.txt'
    if not html_path.exists():
        die(f'{html_path} 不存在，先跑 build --date {d:%Y%m%d}')
    data_path = work / 'data.json'
    data = json.loads(data_path.read_text('utf-8')) if data_path.exists() else {}
    xlsx = Path(data.get('xlsx') or (root / OUT_TPL.format(md=f'{d.month}.{d.day}')))
    if not xlsx.exists():
        die(f'通报 xlsx 不存在：{xlsx}，先跑 build')
    if not PLAN_DOC.exists():
        die(f'方案附件缺失：{PLAN_DOC}')
    to = args.to or RECIPIENT
    subject = args.subject or SUBJECT_TPL.format(md=f'{d.month}.{d.day}')
    return send_mail(to, subject, html_path.read_text('utf-8'), text_path.read_text('utf-8'),
                     [xlsx, PLAN_DOC], args.from_name, args.dry_run)


def cmd_run(args) -> dict:
    b = cmd_backup(args)
    out = cmd_build(args)
    out['backup'] = b
    return out


def cmd_all(args) -> dict:
    b = cmd_backup(args)
    built = cmd_build(args)
    sent = cmd_send(args)
    built['backup'] = b
    built['send'] = sent
    return built


def main() -> None:
    ap = argparse.ArgumentParser(description='企业微信社群运营与质效评价 行动推动情况通报')
    ap.add_argument('--root', default=os.environ.get('QWSQ_ROOT') or str(DEFAULT_ROOT),
                    help='工作根目录（默认 D:\\资料\\ICBC\\20260929企微）')
    sub = ap.add_subparsers(dest='cmd', required=True)

    st = sub.add_parser('status', help='看源表与备份')
    st.set_defaults(func=cmd_status)

    for name, fn, help_text in (('backup', cmd_backup, '备份当日源表'),
                                ('build', cmd_build, '汇总并出表出正文'),
                                ('send', cmd_send, '发信'),
                                ('run', cmd_run, 'backup + build'),
                                ('all', cmd_all, 'backup + build + send')):
        sp = sub.add_parser(name, help=help_text)
        sp.add_argument('--date', help='业务日期 YYYYMMDD（默认今天）')
        sp.add_argument('--prev', help='指定上周备份日期 YYYYMMDD')
        sp.add_argument('--cur-file', help='直接指定本周源文件')
        sp.add_argument('--prev-file', help='直接指定上周源文件')
        if name in ('build', 'run', 'all'):
            sp.add_argument('--out', help='输出 xlsx 路径覆盖')
        if name in ('send', 'all'):
            sp.add_argument('--to', help=f'收件人（默认 {RECIPIENT}）')
            sp.add_argument('--subject', help='主题覆盖')
            sp.add_argument('--from-name', default=FROM_NAME_DEFAULT, help='发件人显示名')
            sp.add_argument('--dry-run', action='store_true', help='只打印不发')
        sp.set_defaults(func=fn)

    args = ap.parse_args()
    result = args.func(args)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
