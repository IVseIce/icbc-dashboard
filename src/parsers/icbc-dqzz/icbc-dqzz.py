#!/home/isee/.venvs/data-tools/bin/python
# -*- coding: utf-8 -*-
"""icbc-dqzz —— 阜阳分行「定期存款增长」通报：取信 → 识别两张月表 → 出通报工作簿 → 回信。

源数据 = 两张《中国工商银行储蓄存款产品余额-分品种、期限时点报表》
（一张上月快照、一张本月快照；邮件里可能是两个 xlsx，也可能是一个含两张月表的 xls 工作簿）。

通报表（sheet「通报」）每一格的口径，完全照抄原表公式：

  B 定期余额   = Σ 本月!{J,L,N,P,R,T,V,X,Z,AB,AD,AF,AH}  ← 「定期」= 整存整取 6 个期限 + 大额存单 7 个期限，
  E 上月余额   = Σ 上月!{同一批 13 列}                     活期/定活两便/通知/零存整取/教育储蓄/存本取息都不算
  C 当月增量   = Σ 本月!比年初列 − Σ 上月!比年初列          （比年初 = 余额 − 年初，两月之差即当月增量）
  D 当月增幅   = C / E                                     （0.00%）
  合计行       = 8 个支行 SUM

支行归属（网点行按名称关键词认行，运行管理部不计入）：
  颍州 = 城建 + 文峰 + 颍州营业室     颍泉 = 颍泉营业室 + 牡丹     颍东 = 颍东营业室 + 汇通
  界首 / 临泉 / 阜南 / 颍上 / 太和 = 各自的营业室

输出 = 三张 sheet 的活公式工作簿（通报 + 上月 + 本月），结构与同事的
「M月D日月定期存款增长.xls」一致：月表里生成 20~27 行的支行汇总块，
通报 sheet 用公式引用它，Excel/WPS 打开即自动重算。

用法
----
  icbc-dqzz.py find   [--keyword 定期] [--date YYYY-MM-DD] [--days 3]
  icbc-dqzz.py fetch  --uid U [--uid U ...] [--root DIR] [--workdir DIR]
  icbc-dqzz.py build  --workdir DIR [--out-name NAME]
  icbc-dqzz.py send   --workdir DIR [--dry-run]
  icbc-dqzz.py run    --uid U [--uid U ...]      # fetch + build，不发信
"""
from __future__ import annotations

import argparse
import email
import imaplib
import io
import json
import os
import re
import smtplib
import socket
import ssl
import sys
import tempfile
import warnings
from copy import copy
from datetime import date, datetime, timedelta
from email.header import make_header, decode_header
from email.message import EmailMessage
from email.utils import formatdate, parsedate_to_datetime, parseaddr
from pathlib import Path
from zoneinfo import ZoneInfo

try:
    import openpyxl
    import xlrd
    import yaml
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    print(json.dumps({'ok': False, 'error': f'缺少依赖: {exc}；请用 /home/isee/.venvs/data-tools/bin/python 运行'},
                     ensure_ascii=False))
    raise SystemExit(2)

warnings.filterwarnings('ignore')

TZ = ZoneInfo('Asia/Shanghai')
DSH_HOME = Path(os.environ.get('DSH_HOME') or (Path.home() / '.dsh'))
SKILL_DIR = Path(__file__).resolve().parent
TEMPLATE = SKILL_DIR / 'assets' / 'template.xlsx'
DEFAULT_ROOT = Path('/mnt/d/资料/ICBC/automations')
WORKDIR_SUFFIX = '_定期存款增长'
KEYWORD = '定期'                       # 邮件主题关键词（find 用）
NAME_KEYWORDS = ('分品种', '时点报表', '定期', '储蓄存款')   # 附件名预筛
SHEET_HINT = ('分品种', '时点报表')     # 表 A1 必须含这两段
REPORT_SHEET = '通报'

# 定期 13 列（余额列；比年初列 = 余额列 + 1）
DQ_COLS = ('J', 'L', 'N', 'P', 'R', 'T', 'V', 'X', 'Z', 'AB', 'AD', 'AF', 'AH')
AGG_FIRST_COL, AGG_LAST_COL = 10, 35    # J..AI：月表里生成汇总块的列范围
AGG_LABEL_COL = 9                       # I：汇总块的支行名列
AGG_FIRST_ROW = 20                      # 汇总块起始行（与原文件一致）
BRANCH_ORDER = ('颍州', '颍泉', '颍东', '界首', '临泉', '阜南', '颍上', '太和')
BRANCH_MAP = {
    '颍州': ('城建', '文峰', '颍州'),
    '颍泉': ('颍泉', '牡丹'),
    '颍东': ('颍东', '汇通'),
    '界首': ('界首',),
    '临泉': ('临泉',),
    '阜南': ('阜南',),
    '颍上': ('颍上',),
    '太和': ('太和',),
}
EXCLUDE_KEYWORDS = ('运行管理部',)
MIN_NETS = 10                            # 一张表至少认出这么多网点才算可用
MAX_ATT_BYTES = 20 * 1024 * 1024
MAX_FILES_PER_SCOPE = 40

PROVIDERS = {
    '163': ('imap.163.com', 993, 'smtp.163.com', 465),
    '126': ('imap.126.com', 993, 'smtp.126.com', 465),
    'qq': ('imap.qq.com', 993, 'smtp.qq.com', 465),
    'sina': ('imap.sina.com.cn', 993, 'smtp.sina.com', 465),
    'aliyun': ('imap.aliyun.com', 993, 'smtp.aliyun.com', 465),
    'gmail': ('imap.gmail.com', 993, 'smtp.gmail.com', 465),
    'outlook': ('outlook.office365.com', 993, 'smtp.office365.com', 587),
    'icloud': ('imap.mail.me.com', 993, 'smtp.mail.me.com', 587),
}


# ================================================================ 邮箱账号
def _read_row_config() -> dict:
    out: dict = {}
    for cfg in sorted((DSH_HOME / 'profiles').glob('*/cordis.patch.yml')):
        try:
            data = yaml.safe_load(cfg.read_text('utf-8')) or []
        except Exception:
            continue
        if not isinstance(data, list):
            continue
        for entry in data:
            if isinstance(entry, dict) and entry.get('id') == 'tool-email':
                conf = entry.get('config')
                if isinstance(conf, dict):
                    out.update(conf)
    return out


def _read_page_config() -> dict:
    try:
        data = yaml.safe_load((DSH_HOME / 'settings.yaml').read_text('utf-8')) or {}
    except Exception:
        return {}
    ns = data.get('dsh-email')
    return ns if isinstance(ns, dict) else {}


def _merge_config(row: dict, page: dict) -> dict:
    cfg = dict(row)
    for key in ('provider', 'user', 'password', 'inboxFolder', 'downloadDir'):
        value = page.get(key)
        if isinstance(value, str) and value.strip():
            cfg[key] = value.strip()
    return cfg


def _pick_account(cfg: dict) -> dict:
    entries, default = None, None
    raw_yaml = cfg.get('accountsYaml') if isinstance(cfg.get('accountsYaml'), str) else ''
    if (raw_yaml or '').strip():
        doc = yaml.safe_load(raw_yaml) or {}
        if isinstance(doc, dict):
            doc = dict(doc)
            default = doc.pop('defaultAccount', None)
            entries = doc
    elif isinstance(cfg.get('accounts'), dict) and cfg['accounts']:
        entries = cfg['accounts']
    if not entries:
        return cfg
    name = cfg.get('defaultAccount') or default or ('default' if 'default' in entries else sorted(entries)[0])
    account = entries.get(name)
    if not isinstance(account, dict):
        raise RuntimeError(f'accountsYaml 里没有默认账号 "{name}"')
    merged = {k: v for k, v in account.items() if v not in (None, '')}
    for key in ('provider', 'user', 'password', 'inboxFolder', 'imap', 'smtp'):
        if key not in merged and cfg.get(key):
            merged[key] = cfg[key]
    return merged


def _env_password():
    pw = os.environ.get('DSH_EMAIL_PASSWORD')
    if pw and pw.strip():
        return pw.strip(), 'DSH_EMAIL_PASSWORD 环境变量'
    for pid in os.listdir('/proc'):
        if not pid.isdigit():
            continue
        try:
            cmd = Path(f'/proc/{pid}/cmdline').read_bytes().replace(b'\0', b' ').decode('utf-8', 'ignore')
            if 'dsh' not in cmd or 'web' not in cmd:
                continue
            for kv in Path(f'/proc/{pid}/environ').read_bytes().split(b'\0'):
                if kv.startswith(b'DSH_EMAIL_PASSWORD='):
                    return kv.split(b'=', 1)[1].decode('utf-8', 'ignore').strip(), '运行中 dsh web 的环境'
        except (OSError, PermissionError):
            continue
    return None


def resolve_account() -> dict:
    cfg = _pick_account(_merge_config(_read_row_config(), _read_page_config()))
    provider = (cfg.get('provider') or '163').strip() or '163'
    imap_host, imap_port, smtp_host, smtp_port = PROVIDERS.get(provider, PROVIDERS['163'])
    imap_conf = cfg.get('imap') if isinstance(cfg.get('imap'), dict) else {}
    smtp_conf = cfg.get('smtp') if isinstance(cfg.get('smtp'), dict) else {}
    pwd = cfg.get('password')
    if isinstance(pwd, str) and pwd.strip():
        password, source = pwd.strip(), '插件配置（设置页/行配置）'
    else:
        found = _env_password()
        if not found:
            raise RuntimeError('取不到邮箱授权码：请在邮件设置页填写，或设置 DSH_EMAIL_PASSWORD')
        password, source = found
    return {
        'provider': provider,
        'user': str(cfg.get('user') or '').strip(),
        'imap_host': str(imap_conf.get('host') or '').strip() or imap_host,
        'imap_port': int(imap_conf.get('port') or imap_port),
        'smtp_host': str(smtp_conf.get('host') or '').strip() or smtp_host,
        'smtp_port': int(smtp_conf.get('port') or smtp_port),
        'password': password,
        'password_source': source,
        'inbox': str(cfg.get('inboxFolder') or 'INBOX'),
    }


def force_ipv4() -> None:
    original = socket.getaddrinfo

    def ipv4_only(host, port, family=0, type=0, proto=0, flags=0):
        return original(host, port, socket.AF_INET, type, proto, flags)

    socket.getaddrinfo = ipv4_only


def connect(acct: dict) -> imaplib.IMAP4_SSL:
    force_ipv4()
    imap = imaplib.IMAP4_SSL(acct['imap_host'], acct['imap_port'],
                             ssl_context=ssl.create_default_context(), timeout=60)
    imap.login(acct['user'], acct['password'])
    try:
        imap.xatom('ID', '("name" "dsh-icbc-dqzz" "version" "1.0" "vendor" "dsh")')
    except Exception:
        pass
    return imap


def _decode(value) -> str:
    if value is None:
        return ''
    try:
        return str(make_header(decode_header(str(value))))
    except Exception:
        return str(value)


def dedup_addresses(addrs) -> list:
    seen, out = set(), []
    for a in addrs:
        a = (a or '').strip()
        if not a or a.lower() in seen:
            continue
        seen.add(a.lower())
        out.append(a)
    return out


# ================================================================ 取信
def fetch_message(imap, uid: str) -> dict:
    typ, body = imap.uid('FETCH', uid, '(BODY.PEEK[])')
    if typ != 'OK' or not body or not isinstance(body[0], tuple):
        raise RuntimeError(f'UID {uid} 取信失败')
    msg = email.message_from_bytes(body[0][1])
    try:
        received = parsedate_to_datetime(msg['Date']).astimezone(TZ)
    except Exception:
        received = None
    atts = []
    for part in msg.walk():
        filename = _decode(part.get_filename())
        if not filename:
            continue
        payload = part.get_payload(decode=True)
        if payload is None:
            continue
        atts.append({'filename': filename, 'blob': payload})
    return {
        'uid': uid,
        'subject': _decode(msg['Subject']),
        'from': _decode(msg['From']),
        'message_id': (msg['Message-ID'] or '').strip(),
        'received': received,
        'received_date': received.date().isoformat() if received else None,
        'attachments': atts,
    }


_MONTHS = ('Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec')


def _imap_date(d: date) -> str:
    return f'{d.day:02d}-{_MONTHS[d.month - 1]}-{d.year}'


def list_recent(imap, since: date, cap: int = 400) -> list:
    """列出 SINCE 之后的邮件抬头。必须用 UID SEARCH：普通 SEARCH 回的是序号。"""
    typ, data = imap.uid('SEARCH', None, f'(SINCE {_imap_date(since)})')
    if typ != 'OK':
        raise RuntimeError(f'IMAP SEARCH 失败：{typ}')
    uids = (data[0] or b'').split()[-cap:]
    out = []
    for u in uids:
        uid = u.decode()
        try:
            typ, d = imap.uid('FETCH', uid, '(BODY.PEEK[HEADER.FIELDS (SUBJECT DATE FROM)])')
        except Exception:
            continue
        if typ != 'OK' or not d or not isinstance(d[0], tuple):
            continue
        msg = email.message_from_bytes(d[0][1])
        try:
            received = parsedate_to_datetime(msg['Date']).astimezone(TZ)
        except Exception:
            received = None
        out.append({'uid': uid, 'subject': _decode(msg['Subject']), 'from': _decode(msg['From']),
                    'received': received,
                    'received_date': received.date().isoformat() if received else None})
    out.sort(key=lambda x: x['received'] or datetime.min.replace(tzinfo=TZ))
    return out


# ================================================================ 数值与表格读取
def _s(v) -> str:
    if v is None:
        return ''
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def _num(v) -> float:
    """源表里的数字有的带千分位、有的混着 '--'，统一成 float；认不出的按 0。"""
    if v is None or isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(',', '').replace('\t', '').replace(' ', '')
    if s in ('', '-', '--', '---', '—', 'null', 'None'):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def _norm_value(v):
    """月表落盘前统一：千分位字符串 → 数字，'--' → 0。"""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return v
    s = str(v).strip()
    if s == '':
        return None
    t = s.replace(',', '').replace('\t', '')
    if t in ('-', '--', '---', '—'):
        return 0
    try:
        return float(t)
    except ValueError:
        return v


def _parse_cn_date(text) -> date | None:
    m = re.search(r'(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日', str(text or ''))
    if not m:
        return None
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


def _is_xls(blob: bytes) -> bool:
    return blob[:8] == b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'


def _looks_like_report(rows: list):
    """是不是《…分品种、期限时点报表》：认 A1 标题 + 合计行。"""
    if not rows or not rows[0]:
        return None
    a1 = _s(rows[0][0] if rows[0] else '')
    if not all(k in a1 for k in SHEET_HINT):
        return None
    d = _parse_cn_date(a1)
    if not d:
        return None
    ncols = max((len(r) for r in rows if r), default=0)
    total_row, max_row = None, 0
    for i, r in enumerate(rows):
        name = _s(r[0]) if r else ''
        if name == '合计':
            total_row = i + 1
        if name and i + 1 >= 5 and name != '合计':
            max_row = i + 1
    if not total_row or max_row <= total_row:
        return None
    return {'date': d, 'max_row': max_row, 'ncols': ncols, 'total_row': total_row}


def _norm_rows(rows: list) -> list:
    return [[_norm_value(v) for v in r] for r in rows]


def _rows_of_xlsx(blob: bytes) -> list:
    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    try:
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _rows_of_xls(blob: bytes) -> list:
    with tempfile.NamedTemporaryFile(suffix='.xls', delete=False) as fh:
        fh.write(blob)
        path = fh.name
    try:
        book = xlrd.open_workbook(path, formatting_info=True)
        out = []
        for sh in book.sheets():
            rows = [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
            out.append((sh.name, rows))
        return out
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass


def read_tables(blob: bytes, filename: str, uid: str = '', received=None) -> list:
    """把一个表格附件里所有「分品种、期限时点报表」sheet 读成候选项。"""
    try:
        sheets = _rows_of_xls(blob) if _is_xls(blob) else _rows_of_xlsx(blob)
    except Exception:
        return []
    out = []
    for title, rows in sheets:
        meta = _looks_like_report(rows)
        if not meta:
            continue
        rows = _norm_rows(rows)
        nets, extra, excluded = _map_nets(rows, meta['total_row'], meta['max_row'])
        out.append({
            'kind': 'xls' if _is_xls(blob) else 'xlsx',
            'file': filename, 'sheet': title, 'uid': uid,
            'received': received, 'received_date': received.date().isoformat() if received else None,
            'date': meta['date'], 'total_row': meta['total_row'], 'max_row': meta['max_row'],
            'ncols': meta['ncols'], 'rows': rows, 'nets': nets, 'extra_rows': extra,
            'excluded': excluded,
        })
    return out


def _match_branch(name: str):
    s = _s(name)
    if not s:
        return None
    for kw in EXCLUDE_KEYWORDS:
        if kw in s:
            return None
    for branch, kws in BRANCH_MAP.items():
        for kw in kws:
            if kw in s:
                return branch
    return None


def _map_nets(rows: list, total_row: int, max_row: int):
    """→ ({支行: 行号}, 认不出归属的行名, 被排除的行名)"""
    nets: dict = {}
    extra, excluded = [], []
    for r in range(total_row + 1, max_row + 1):
        if r - 1 >= len(rows):
            break
        name = _s(rows[r - 1][0] if rows[r - 1] else '')
        if not name or name == '合计':
            continue
        if any(kw in name for kw in EXCLUDE_KEYWORDS):
            excluded.append(name)
            continue
        branch = _match_branch(name)
        if not branch:
            extra.append(name)
            continue
        nets.setdefault(branch, []).append(r)
    return nets, extra, excluded


# ================================================================ 选表
def _table_key(t: dict) -> str:
    return f'{t["file"]}#{t["sheet"]}'


def _table_brief(t: dict) -> dict:
    return {'file': t['file'], 'sheet': t['sheet'], 'date': t['date'].isoformat(),
            'uid': t.get('uid'), 'nets': sorted(t.get('nets') or {}),
            'net_count': sum(len(v) for v in (t.get('nets') or {}).values()),
            'received_date': t.get('received_date')}


def _rank(t: dict):
    """同一日期多张表时的取舍：网点认得多的优先，其次后来者（新邮件）优先。"""
    return (sum(len(v) for v in (t.get('nets') or {}).values()), t.get('received') or datetime.min.replace(tzinfo=TZ))


def _month_end(d: date) -> bool:
    return (d + timedelta(days=1)).day == 1


def resolve_tables(cands: list) -> dict:
    """按 A1 日期选出「本月（最新快照）」与「上月（之前最近一张）」。"""
    problems, notices = [], []
    # 同一个文件同一张表可能被两封邮件各带一份，按 (日期, 文件, 表名) 去重
    uniq: dict = {}
    for t in cands:
        key = (t['date'].isoformat(), t['file'], t['sheet'])
        stamp = t.get('received') or datetime.min.replace(tzinfo=TZ)
        if key not in uniq or stamp > (uniq[key].get('received') or datetime.min.replace(tzinfo=TZ)):
            uniq[key] = t
    cands = list(uniq.values())
    good = [t for t in cands if sum(len(v) for v in (t.get('nets') or {}).values()) >= MIN_NETS]
    weak = [t for t in cands if t not in good]
    if weak:
        notices.append('这些表认出的网点太少（<%d 个），已跳过：%s'
                       % (MIN_NETS, '、'.join(_table_key(t) for t in weak)))
    if not good:
        return {'ok': False, 'prev': None, 'cur': None, 'notices': notices,
                'problems': ['没有一张可用的《分品种、期限时点报表》（认出的网点不足 %d 个）' % MIN_NETS]}

    dated = sorted(good, key=lambda t: (t['date'], _rank(t)))
    cur = dated[-1]
    before = [t for t in dated if t['date'] < cur['date']]
    prev = before[-1] if before else None

    if prev is None:
        problems.append('只找到一张月表（最新快照 %s），缺上月快照，无法算当月增量'
                        % cur['date'].isoformat())
    elif not _month_end(prev['date']):
        notices.append('上月取的是 %s 的快照（不是月末）：%s。如需改用月末表请指定邮件'
                       % (prev['date'].isoformat(), _table_key(prev)))

    same_day = [t for t in dated if prev and t['date'] == prev['date']]
    if len(same_day) > 1:
        notices.append('同一天有 %d 张表，上月取了：%s；其余（内容应相同）：%s'
                       % (len(same_day), _table_key(prev),
                          '、'.join(_table_key(t) for t in same_day if t is not prev)))
    same_cur = [t for t in dated if t['date'] == cur['date']]
    if len(same_cur) > 1:
        notices.append('本月同一天有 %d 张表，取了：%s；其余（内容应相同）：%s'
                       % (len(same_cur), _table_key(cur),
                          '、'.join(_table_key(t) for t in same_cur if t is not cur)))
    others = [t for t in dated if t is not cur and t is not prev]
    if others:
        notices.append('邮件里另有 %d 张更早的月表未使用：%s'
                       % (len(others), '、'.join(f'{_table_key(t)}（{t["date"].isoformat()}）' for t in others[:9])))

    for t in (prev, cur):
        if not t:
            continue
        miss = [b for b in BRANCH_ORDER if b not in (t.get('nets') or {})]
        if miss:
            notices.append('%s：源表里没认出这些支行，按 0 算：%s' % (_table_key(t), '、'.join(miss)))
        if t.get('extra_rows'):
            notices.append('%s：这些行不属于 8 个支行，已忽略：%s'
                           % (_table_key(t), '、'.join(t['extra_rows'])))
        if t.get('excluded'):
            notices.append('%s：不计入通报的行：%s' % (_table_key(t), '、'.join(t['excluded'])))

    return {'ok': not problems, 'prev': prev, 'cur': cur,
            'problems': problems, 'notices': notices}


# ================================================================ 计算
def _col_vals(rows: list, r: int, cols, off: int = 0) -> list:
    row = rows[r - 1] if r - 1 < len(rows) else []
    out = []
    for col in cols:
        idx = _col_index(col) - 1 + off
        out.append(_num(row[idx]) if idx < len(row) else 0.0)
    return out


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n


def compute(prev_t: dict, cur_t: dict) -> dict:
    """按原表公式算出 8 个支行 + 合计。"""
    branches: dict = {}
    warnings: list = []
    for branch in BRANCH_ORDER:
        prows = (prev_t.get('nets') or {}).get(branch) or []
        crows = (cur_t.get('nets') or {}).get(branch) or []
        prev_bal = sum(sum(_col_vals(prev_t['rows'], r, DQ_COLS)) for r in prows)
        cur_bal = sum(sum(_col_vals(cur_t['rows'], r, DQ_COLS)) for r in crows)
        inc = 0.0
        for r in crows:
            inc += sum(_col_vals(cur_t['rows'], r, DQ_COLS, off=1))
        for r in prows:
            inc -= sum(_col_vals(prev_t['rows'], r, DQ_COLS, off=1))
        if not prows:
            warnings.append(f'「{branch}」在上月表里没有对应网点行，上月余额按 0 算')
        if not crows:
            warnings.append(f'「{branch}」在本月表里没有对应网点行，本月余额按 0 算')
        branches[branch] = {
            'prev_bal': round(prev_bal, 2), 'cur_bal': round(cur_bal, 2),
            'inc': round(inc, 2),
            'rate': (inc / prev_bal) if prev_bal else None,
            'prev_rows': prows, 'cur_rows': crows,
            'prev_nets': len(prows), 'cur_nets': len(crows),
        }
    total = {
        'prev_bal': round(sum(b['prev_bal'] for b in branches.values()), 2),
        'cur_bal': round(sum(b['cur_bal'] for b in branches.values()), 2),
        'inc': round(sum(b['inc'] for b in branches.values()), 2),
    }
    total['rate'] = (total['inc'] / total['prev_bal']) if total['prev_bal'] else None
    return {'branches': branches, 'total': total, 'warnings': warnings}


# ================================================================ 月表样式
_XLS_HALIGN = {0: None, 1: 'left', 2: 'center', 3: 'right', 4: 'justify', 5: 'fill', 6: 'centerContinuous',
               7: 'distributed'}
_XLS_VALIGN = {0: 'top', 1: 'center', 2: 'bottom', 3: 'justify', 4: 'distributed'}
_XLS_LINE = {0: None, 1: 'thin', 2: 'medium', 3: 'dashed', 4: 'dotted', 5: 'thick', 6: 'double', 7: 'hair'}
_XLS_COLOUR = {0: 'FF000000', 8: 'FF000000', 9: 'FFFFFFFF', 10: 'FFFF0000', 11: 'FF00B050', 22: 'FFC00000'}


def _xls_style(book, sh, r: int, c: int):
    """(font, border, alignment, number_format, fill) —— 从 xls 的 XF 记录还原。"""
    xf = book.xf_list[sh.cell_xf_index(r, c)]
    f = book.font_list[xf.font_index]
    font = Font(name=f.name or '宋体', size=(f.height / 20.0) if f.height else 11,
                bold=bool(f.bold) or int(f.weight or 0) >= 700, italic=bool(f.italic),
                color=_XLS_COLOUR.get(f.colour_index))
    bd = xf.border
    border = Border(left=Side(style=_XLS_LINE.get(bd.left_line_style)),
                    right=Side(style=_XLS_LINE.get(bd.right_line_style)),
                    top=Side(style=_XLS_LINE.get(bd.top_line_style)),
                    bottom=Side(style=_XLS_LINE.get(bd.bottom_line_style)))
    al = xf.alignment
    alignment = Alignment(horizontal=_XLS_HALIGN.get(al.hor_align), vertical=_XLS_VALIGN.get(al.vert_align),
                          wrap_text=bool(al.text_wrapped))
    fmt = book.format_map[xf.format_key].format_str or 'General'
    fill = None
    if xf.background.fill_pattern not in (0, None) and xf.background.pattern_colour_index not in (64, 65):
        colour = _XLS_COLOUR.get(xf.background.pattern_colour_index)
        if colour:
            fill = PatternFill(fill_type='solid', fgColor=colour)
    return font, border, alignment, fmt, fill


def _copy_from_xlsx(blob: bytes, sheet: str, dst_ws, max_row: int, ncols: int) -> None:
    wb = openpyxl.load_workbook(io.BytesIO(blob))          # 带格式（数据区没有公式，照抄无妨）
    try:
        ws = wb[sheet]
        for r in range(1, min(max_row, ws.max_row) + 1):
            for c in range(1, min(ncols, ws.max_column) + 1):
                src, dst = ws.cell(row=r, column=c), dst_ws.cell(row=r, column=c)
                if src.has_style:
                    dst.font = copy(src.font)
                    dst.border = copy(src.border)
                    dst.fill = copy(src.fill)
                    dst.alignment = copy(src.alignment)
                    dst.protection = copy(src.protection)
                    dst.number_format = src.number_format
        for key, dim in ws.column_dimensions.items():
            if dim.width:
                dst_ws.column_dimensions[key].width = dim.width
            if dim.hidden:
                dst_ws.column_dimensions[key].hidden = True
        for key, dim in ws.row_dimensions.items():
            try:
                if dim.height and int(key) <= max_row:
                    dst_ws.row_dimensions[key].height = dim.height
            except (TypeError, ValueError):
                pass
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= max_row:
                dst_ws.merge_cells(str(rng))
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _copy_from_xls(blob: bytes, sheet: str, dst_ws, max_row: int, ncols: int) -> None:
    with tempfile.NamedTemporaryFile(suffix='.xls', delete=False) as fh:
        fh.write(blob)
        path = fh.name
    try:
        book = xlrd.open_workbook(path, formatting_info=True)
        sh = book.sheet_by_name(sheet)
        for r in range(min(max_row, sh.nrows)):
            for c in range(min(ncols, sh.ncols)):
                font, border, alignment, fmt, fill = _xls_style(book, sh, r, c)
                dst = dst_ws.cell(row=r + 1, column=c + 1)
                dst.font, dst.border, dst.alignment = font, border, alignment
                dst.number_format = fmt
                if fill:
                    dst.fill = fill
        wanted = {'宋体', 'DejaVu Sans'}
        for key, info in sh.colinfo_map.items():
            if info.width:
                dst_ws.column_dimensions[get_column_letter(key + 1)].width = info.width / 256.0
        for key, info in sh.rowinfo_map.items():
            if info.height and key + 1 <= max_row:
                dst_ws.row_dimensions[key + 1].height = info.height / 20.0
        for rlo, rhi, clo, chi in getattr(sh, 'merged_cells', []):
            if rlo < max_row:
                dst_ws.merge_cells(start_row=rlo + 1, start_column=clo + 1,
                                   end_row=rhi, end_column=chi)
        # xls 里的默认字体常被映射成 DejaVu Sans，统一回宋体
        for row in dst_ws.iter_rows(min_row=1, max_row=max_row, max_col=min(ncols, dst_ws.max_column)):
            for cell in row:
                if cell.font and cell.font.name in wanted and cell.font.name != '宋体':
                    cell.font = Font(name='宋体', size=cell.font.sz, bold=cell.font.b,
                                     italic=cell.font.i, color=cell.font.color)
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass


# ================================================================ 出表
def _sheet_name_for(t: dict, used: set) -> str:
    name = f'{t["date"].month}月'
    if name in used or not re.fullmatch(r'[\w\u4e00-\u9fff]{1,30}', name):
        name = '本月' if not used else '上月'
    used.add(name)
    return name


def generation(prev_t: dict, cur_t: dict):
    """两张月表 → (通报 8 行公式 + 合计公式, 月表汇总块公式)。"""
    def sum_of(sheet: str, col: str, row: int, off: int = 0) -> str:
        letter = _col_offset(col, off)
        return f'{sheet}!{letter}{row}'

    def expr(sheet: str, row: int, off: int = 0) -> str:
        return '+'.join(sum_of(sheet, c, row, off) for c in DQ_COLS)

    report = {}
    for i, branch in enumerate(BRANCH_ORDER):
        r, ar = 3 + i, AGG_FIRST_ROW + i
        report[r] = {
            'B': '=' + expr(cur_t['sheet_name'], ar),
            'C': '=' + expr(cur_t['sheet_name'], ar, off=1) + '-(' + expr(prev_t['sheet_name'], ar, off=1) + ')',
            'D': f'=C{r}/E{r}',
            'E': '=' + expr(prev_t['sheet_name'], ar),
        }
    report['total'] = {'B': '=SUM(B3:B10)', 'C': '=SUM(C3:C10)', 'D': '=C11/E11', 'E': '=SUM(E3:E10)'}

    agg = {}
    for t in (prev_t, cur_t):
        rows_map = {}
        for i, branch in enumerate(BRANCH_ORDER):
            ar = AGG_FIRST_ROW + i
            members = (t.get('nets') or {}).get(branch) or []
            cells = {}
            for c in range(AGG_FIRST_COL, AGG_LAST_COL + 1):
                letter = get_column_letter(c)
                cells[letter] = ('=' + '+'.join(f'{letter}{r}' for r in members)) if members else '=0'
            rows_map[ar] = {'label': branch, 'cells': cells}
        agg[t['sheet_name']] = rows_map
    return report, agg


def _col_offset(letters: str, off: int) -> str:
    return get_column_letter(_col_index(letters) + off)


def build_workbook(prev_t: dict, cur_t: dict, out_path: Path) -> dict:
    if not TEMPLATE.exists():
        raise RuntimeError(f'内置模板缺失：{TEMPLATE}')
    wb = openpyxl.load_workbook(TEMPLATE)
    ws_report = wb[REPORT_SHEET]

    used: set = {REPORT_SHEET}
    for t in (prev_t, cur_t):
        t['sheet_name'] = _sheet_name_for(t, used)
    report, agg = generation(prev_t, cur_t)

    # 1) 两张月表（值 + 格式 + 汇总块）
    for t in (prev_t, cur_t):
        dst = wb.create_sheet(t['sheet_name'])
        rows = t['rows']
        for r in range(1, min(t['max_row'], len(rows)) + 1):
            row = rows[r - 1]
            for c in range(1, min(t['ncols'], len(row)) + 1):
                dst.cell(row=r, column=c, value=row[c - 1])
        try:
            if t['kind'] == 'xlsx':
                _copy_from_xlsx(t['blob'], t['sheet'], dst, t['max_row'], t['ncols'])
            else:
                _copy_from_xls(t['blob'], t['sheet'], dst, t['max_row'], t['ncols'])
        except Exception as exc:
            t.setdefault('style_notes', []).append(f'月表格式复制失败（{type(exc).__name__}: {exc}），已只写数值')
        # 数据区统一成 #,##0.00（与原文件里的月表一致）
        for r in range(t['total_row'], t['max_row'] + 1):
            for c in range(2, min(t['ncols'], dst.max_column) + 1):
                cell = dst.cell(row=r, column=c)
                if isinstance(cell.value, (int, float)):
                    cell.font = copy(cell.font)
                    cell.number_format = '#,##0.00'
        # 汇总块（20~27 行）：与原文件同样的公式
        for ar, info in agg[t['sheet_name']].items():
            lab = dst.cell(row=ar, column=AGG_LABEL_COL, value=info['label'])
            lab.font = Font(name='宋体', size=11)
            lab.alignment = Alignment(horizontal='center', vertical='center')
            for letter, formula in info['cells'].items():
                cell = dst[f'{letter}{ar}']
                cell.value = formula
                cell.font = Font(name='宋体', size=12)
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.number_format = '0_ '

    # 2) 通报 sheet：标题、表头、公式
    title = _report_title(cur_t['date'])
    ws_report['A1'] = title
    ws_report['E2'] = f'{prev_t["date"].month}月定期余额'
    for r, cells in report.items():
        if r == 'total':
            for col, formula in cells.items():
                ws_report[f'{col}11'] = formula
        else:
            for col, formula in cells.items():
                ws_report[f'{col}{r}'] = formula

    wb.calculation.fullCalcOnLoad = True
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {'title': title, 'sheets': wb.sheetnames,
            'report_sheet': REPORT_SHEET, 'out': str(out_path)}


def _report_title(d: date | None) -> str:
    if not d:
        return '阜阳分行定期增长情况表（单位：万元）'
    return f'{d.month}月{d.day}日阜阳分行定期增长情况表（单位：万元）'


def _title_short(d: date | None) -> str:
    if not d:
        return ''
    return f'{d.month}.{d.day}'


# ================================================================ 文案
def build_text(prev_t: dict, cur_t: dict, calc: dict) -> str:
    total = calc['total']
    lines = [
        f'截至{_report_title(cur_t["date"]).split("阜阳分行")[0]}，'
        f'全行定期存款余额 {total["cur_bal"]:,.2f} 万元，当月增量 {total["inc"]:,.2f} 万元，'
        f'当月增幅 {(total["rate"] or 0) * 100:.2f}%（对照 {prev_t["date"].month}月'
        f'{prev_t["date"].day}日 {total["prev_bal"]:,.2f} 万元）。',
        f'附件为通报表（含{prev_t["date"].month}月、{cur_t["date"].month}月原始数据两张表，公式可核对）。',
    ]
    return '\n'.join(lines) + '\n'


# ================================================================ 检索与归类
def _att_candidates(imap, mails, cache, notices: list) -> list:
    """把若干封邮件的表格附件读成候选项（uid 记录来源邮件）。"""
    out, files, capped = [], 0, False
    for m in mails:
        if capped:
            break
        if m['uid'] not in cache:
            cache[m['uid']] = fetch_message(imap, m['uid'])
        for att in cache[m['uid']]['attachments']:
            if not att['filename'].lower().endswith(('.xls', '.xlsx', '.xlsm', '.et')):
                continue
            if not any(k in att['filename'] for k in NAME_KEYWORDS):
                continue
            if len(att['blob']) > MAX_ATT_BYTES:
                notices.append(f'跳过过大的附件 {att["filename"]}（{len(att["blob"]) // 1024}KB）')
                continue
            if files >= MAX_FILES_PER_SCOPE:
                notices.append(f'本次最多解析 {MAX_FILES_PER_SCOPE} 个附件，其余已跳过')
                capped = True
                break
            files += 1
            out.extend(read_tables(att['blob'], att['filename'], m['uid'], m['received']))
    return out


def _pool_check(imap, mails, cache) -> dict:
    notices: list = []
    try:
        cands = _att_candidates(imap, mails, cache, notices)
    except Exception as exc:
        return {'ok': False, 'problems': [f'{type(exc).__name__}: {exc}'], 'notices': notices,
                'prev': None, 'cur': None}
    got = resolve_tables(cands)
    got['notices'] = notices + list(got.get('notices') or [])
    return got


def _brief(m: dict) -> dict:
    return {'uid': m['uid'], 'subject': m['subject'], 'from': m['from'],
            'received_date': m.get('received_date')}


def _workdir_for(root: Path, received: date) -> Path:
    return root / f'{received.strftime("%Y%m%d")}{WORKDIR_SUFFIX}'


def _strip_tables(got: dict) -> dict:
    def one(t):
        if not t:
            return None
        b = _table_brief(t)
        b['extra_rows'] = t.get('extra_rows')
        b['excluded'] = t.get('excluded')
        return b
    return {'ok': got.get('ok'), 'prev': one(got.get('prev')), 'cur': one(got.get('cur')),
            'problems': got.get('problems') or [], 'notices': got.get('notices') or []}


def cmd_find(args) -> dict:
    today = date.fromisoformat(args.date) if args.date else datetime.now(TZ).date()
    keyword = args.keyword
    acct = resolve_account()
    imap = connect(acct)
    try:
        typ, _ = imap.select(acct['inbox'], readonly=True)
        if typ != 'OK':
            raise RuntimeError(f"SELECT {acct['inbox']} 失败")
        recent = list_recent(imap, today - timedelta(days=args.days))
        if args.date:                       # 补做历史时只看当天及更早，免得抓到之后的邮件
            recent = [m for m in recent if m['received_date'] and m['received_date'] <= today.isoformat()]
        cache: dict = {}
        kw = [m for m in recent if keyword and keyword in (m['subject'] or '')]
        kw_today = [m for m in kw if m['received_date'] == today.isoformat()]

        base = {'ok': True, 'keyword': keyword, 'today': today.isoformat(),
                'inbox': acct['inbox'], 'scanned': len(recent),
                'keyword_hits': [_brief(m) for m in kw],
                'keyword_hits_today': [_brief(m) for m in kw_today]}

        why = ''
        if kw_today:
            got = _pool_check(imap, kw_today, cache)
            if got.get('ok'):
                base.update({'auto': True, 'uids': [m['uid'] for m in kw_today],
                             'proposed': [_brief(m) for m in kw_today],
                             'verification': _strip_tables(got),
                             'reason': '主题含关键词「%s」且为当天邮件，上月/本月两张月表齐全' % keyword})
                return base
            why = '；'.join(got.get('problems') or []) or '两张月表不全'
        else:
            why = (f'今天（{today.isoformat()}）没有主题含「{keyword}」的邮件'
                   if not kw else f'主题含「{keyword}」的邮件都不是今天收到的')

        days = sorted({m['received_date'] for m in recent if m['received_date']}, reverse=True)
        scopes = [
            (f'窗口内主题含「{keyword}」的邮件', kw),
            (f'最近一天（{days[0]}）收到的全部邮件' if days else '',
             [m for m in recent if days and m['received_date'] == days[0]]),
            ('窗口内收到的全部邮件', recent),
        ]
        tried, found = [], None
        for label, mails in scopes:
            if not mails:
                continue
            got = _pool_check(imap, mails, cache)
            picked = []
            for t in (got.get('cur'), got.get('prev')):
                if t and t.get('uid') and t['uid'] not in picked:
                    picked.append(t['uid'])
            tried.append({'scope': label, 'scanned': len(mails), 'picked': picked,
                          'complete': bool(got.get('ok')),
                          'tables': [t['date'].isoformat() for t in (got.get('cur'), got.get('prev')) if t],
                          'problems': got.get('problems') or []})
            if got.get('ok'):
                found = {'scope': label, 'uids': picked, 'verification': got}
                break

        base['auto'] = False
        base['reason'] = why
        base['tried'] = tried
        if found:
            uids = found['uids']
            base['uids'] = uids
            base['proposed'] = [_brief(cache[u]) for u in uids]
            base['verification'] = _strip_tables(found['verification'])
            base['question'] = (f'{why}。按附件内容检索（{found["scope"]}），'
                                f'其中 {len(uids)} 封可以凑齐上月/本月两张月表。是否按这几封继续？'
                                f'（uid：{"、".join(uids)}；以哪封为准也可以直接指定）')
        else:
            base['question'] = (f'{why}，且放宽检索也没能凑齐两张月表。'
                                f'已检索 {len(recent)} 封、{args.days} 天内的邮件。'
                                f'请确认源邮件是否已到，或直接指定邮件。')
        return base
    finally:
        try:
            imap.logout()
        except Exception:
            pass


def cmd_fetch(args) -> dict:
    acct = resolve_account()
    imap = connect(acct)
    try:
        typ, _ = imap.select(acct['inbox'], readonly=True)
        if typ != 'OK':
            raise RuntimeError(f"SELECT {acct['inbox']} 失败")
        msgs = [fetch_message(imap, u) for u in args.uid]
    finally:
        try:
            imap.logout()
        except Exception:
            pass

    dated = [m for m in msgs if m['received_date']]
    if not dated:
        return {'ok': False, 'error': '取不到邮件日期，无法命名工作目录'}
    notices: list = []
    cands = []
    for m in msgs:
        for att in m['attachments']:
            if not att['filename'].lower().endswith(('.xls', '.xlsx', '.xlsm', '.et')):
                continue
            cands.extend(read_tables(att['blob'], att['filename'], m['uid'], m['received']))
    got = resolve_tables(cands)
    got['notices'] = notices + list(got.get('notices') or [])
    main = _pick_mail(msgs, got)
    workdir = Path(args.workdir) if args.workdir else _workdir_for(Path(args.root), main['received'].date())
    src = workdir / 'source'
    for m in msgs:
        for att in m['attachments']:
            if att['filename'].lower().endswith(('.xls', '.xlsx', '.xlsm', '.et')):
                src.mkdir(parents=True, exist_ok=True)
                (src / att['filename']).write_bytes(att['blob'])

    meta = {
        'ok': got['ok'],
        'workdir': str(workdir),
        'source_dir': str(src),
        'received_date': main['received_date'],
        'mails': [{'uid': m['uid'], 'subject': m['subject'], 'from': m['from'],
                   'received_date': m['received_date'],
                   'attachments': [a['filename'] for a in m['attachments']]} for m in msgs],
        'tables': {'prev': _strip_tables(got).get('prev'), 'cur': _strip_tables(got).get('cur')},
        'problems': got.get('problems') or [],
        'notices': got.get('notices') or [],
        'reply': {'to': parseaddr(main['from'])[1], 'subject': main['subject'],
                  'message_id': main['message_id'], 'uid': main['uid']},
    }
    if not got['ok']:
        return {**meta, 'error': '；'.join(got.get('problems') or []) or '缺上月/本月月表'}

    (workdir / 'roles.json').write_text(
        json.dumps({'tables': {'prev': _table_ref(got['prev']), 'cur': _table_ref(got['cur'])},
                    'reply': meta['reply'], 'received_date': meta['received_date'],
                    'notices': meta['notices']}, ensure_ascii=False, indent=1), 'utf-8')
    return meta


def _table_ref(t: dict) -> dict:
    return {'file': t['file'], 'sheet': t['sheet'], 'date': t['date'].isoformat(),
            'uid': t.get('uid'), 'received_date': t.get('received_date')}


def _pick_mail(msgs: list, got: dict) -> dict:
    """回信对象：优先带本月表那封，否则最新一封。"""
    cur = got.get('cur')
    if cur and cur.get('uid'):
        hit = [m for m in msgs if m['uid'] == cur['uid']]
        if hit:
            return hit[0]
    return max([m for m in msgs if m['received_date']], key=lambda m: m['received'])


def _load_tables(workdir: Path, source=None) -> dict:
    meta_p = workdir / 'roles.json'
    if not meta_p.exists():
        raise RuntimeError(f'缺少 {meta_p}，请先执行 fetch')
    meta = json.loads(meta_p.read_text('utf-8'))
    src = Path(source) if source else (workdir / 'source')
    out = {}
    for role in ('prev', 'cur'):
        ref = meta['tables'][role]
        blob = (src / ref['file']).read_bytes()
        cands = [t for t in read_tables(blob, ref['file']) if t['sheet'] == ref['sheet']]
        if not cands:
            raise RuntimeError(f'{role} 月表读不出来：{ref["file"]}#{ref["sheet"]}')
        t = cands[0]
        t['blob'] = blob
        out[role] = t
    return {'tables': out, 'reply': meta.get('reply', {}),
            'received_date': meta.get('received_date'), 'notices': meta.get('notices', [])}


# ================================================================ 子命令
def cmd_build(args) -> dict:
    workdir = Path(args.workdir)
    loaded = _load_tables(workdir, getattr(args, 'source', None))
    prev_t, cur_t = loaded['tables']['prev'], loaded['tables']['cur']

    calc = compute(prev_t, cur_t)
    name = args.out_name or f'定期存款增长（{_title_short(cur_t["date"])}）.xlsx'
    out_path = workdir / 'out' / name
    if out_path.parent.is_dir():
        for p in out_path.parent.glob('*.xlsx'):
            if p.name != name:
                try:
                    p.unlink()
                except OSError:
                    pass
    built = build_workbook(prev_t, cur_t, out_path)

    body = build_text(prev_t, cur_t, calc)
    (workdir / '文案.txt').write_text(body, 'utf-8')
    notices = list(dict.fromkeys(list(loaded.get('notices') or [])
                                 + [f'{cur_t["file"]}#{cur_t["sheet"]} → sheet「{cur_t["sheet_name"]}」',
                                    f'{prev_t["file"]}#{prev_t["sheet"]} → sheet「{prev_t["sheet_name"]}」']))
    subtotal = {b: {k: v for k, v in calc['branches'][b].items()
                    if k in ('prev_bal', 'cur_bal', 'inc', 'rate')} for b in BRANCH_ORDER}
    (workdir / 'data.json').write_text(json.dumps(
        {'title': built['title'], 'prev_date': prev_t['date'].isoformat(),
         'cur_date': cur_t['date'].isoformat(), 'branches': subtotal, 'total': calc['total'],
         'warnings': calc['warnings']}, ensure_ascii=False, indent=1), 'utf-8')
    return {'ok': True, 'workdir': str(workdir), 'report': str(out_path), 'report_name': name,
            'title': built['title'], 'sheets': built['sheets'],
            'prev_date': prev_t['date'].isoformat(), 'cur_date': cur_t['date'].isoformat(),
            'branches': subtotal, 'total': calc['total'],
            'warnings': calc['warnings'], 'notices': notices,
            'text_file': str(workdir / '文案.txt'), 'text': body, 'reply': loaded['reply']}


def _reply_subject(d: date | None) -> str:
    short = _title_short(d)
    return f'定期存款增长（{short}）' if short else '定期存款增长'


def cmd_send(args) -> dict:
    workdir = Path(args.workdir)
    loaded = _load_tables(workdir)
    reply = loaded['reply']
    xlsx = sorted((workdir / 'out').glob('*.xlsx'))
    if not xlsx:
        return {'ok': False, 'error': f'{workdir}/out 下没有 xlsx，请先 build'}
    report = xlsx[0]
    txt = workdir / '文案.txt'
    body = txt.read_text('utf-8') if txt.exists() else ''
    to = (args.to or reply.get('to') or '').strip()
    if not to:
        return {'ok': False, 'error': '没有收件人（源邮件没有解析出发件人地址）'}
    subject = args.subject or _reply_subject(loaded['tables']['cur']['date'])

    acct = resolve_account()
    msg = EmailMessage()
    msg['From'] = acct['user']
    msg['To'] = to
    msg['Subject'] = subject
    msg['Date'] = formatdate(localtime=True)
    if reply.get('message_id'):
        msg['In-Reply-To'] = reply['message_id']
        msg['References'] = reply['message_id']
    msg.set_content(body)
    msg.add_attachment(report.read_bytes(), maintype='application',
                       subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                       filename=report.name)
    if args.dry_run:
        return {'ok': True, 'dry_run': True, 'to': to, 'subject': subject,
                'attachment': report.name, 'bytes': len(report.read_bytes()),
                'smtp': f"{acct['smtp_host']}:{acct['smtp_port']}", 'body': body}

    force_ipv4()
    ctx = ssl.create_default_context()
    if acct['smtp_port'] == 465:
        server = smtplib.SMTP_SSL(acct['smtp_host'], acct['smtp_port'], context=ctx, timeout=60)
    else:
        server = smtplib.SMTP(acct['smtp_host'], acct['smtp_port'], timeout=60)
        server.starttls(context=ctx)
    try:
        server.login(acct['user'], acct['password'])
        server.send_message(msg)
    finally:
        try:
            server.quit()
        except Exception:
            pass
    return {'ok': True, 'sent': True, 'to': to, 'subject': subject, 'attachment': report.name,
            'smtp': f"{acct['smtp_host']}:{acct['smtp_port']}"}


def cmd_run(args) -> dict:
    got = cmd_fetch(args)
    if not got.get('ok'):
        return got
    out = cmd_build(argparse.Namespace(workdir=got['workdir'], source=None, out_name=None))
    for k in ('workdir', 'received_date', 'mails'):
        out.setdefault(k, got.get(k))
    out['notices'] = list(dict.fromkeys(list(got.get('notices') or []) + list(out.get('notices') or [])))
    out['fetch'] = got
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description='阜阳分行定期存款增长通报：取信/出表/回信')
    sub = ap.add_subparsers(dest='cmd', required=True)

    d = sub.add_parser('find', help='按约定找源邮件：主题含关键词且当天收到则直接执行，否则报告并请示')
    d.add_argument('--keyword', default=KEYWORD, help=f'主题关键词，默认「{KEYWORD}」')
    d.add_argument('--date', default=None, help='把哪天当作"今天"（YYYY-MM-DD）')
    d.add_argument('--days', type=int, default=3, help='放宽检索的回看天数，默认 3')
    d.set_defaults(func=cmd_find)

    f = sub.add_parser('fetch', help='按 uid 取信并把附件存到工作目录')
    f.add_argument('--uid', action='append', required=True)
    f.add_argument('--root', default=str(DEFAULT_ROOT), help='工作目录根，默认 /mnt/d/资料/ICBC/automations')
    f.add_argument('--workdir', default=None)
    f.set_defaults(func=cmd_fetch)

    b = sub.add_parser('build', help='识出上月/本月两张月表并出通报工作簿（活公式）')
    b.add_argument('--workdir', required=True)
    b.add_argument('--source', default=None)
    b.add_argument('--out-name', default=None, help='输出文件名，默认 定期存款增长（M.D）.xlsx')
    b.set_defaults(func=cmd_build)

    s = sub.add_parser('send', help='把通报回复给原发件人')
    s.add_argument('--workdir', required=True)
    s.add_argument('--to', default=None)
    s.add_argument('--subject', default=None, help='覆盖回信主题，默认「定期存款增长（M.D）」')
    s.add_argument('--dry-run', action='store_true')
    s.set_defaults(func=cmd_send)

    r = sub.add_parser('run', help='fetch + build（不发信）')
    r.add_argument('--uid', action='append', required=True)
    r.add_argument('--root', default=str(DEFAULT_ROOT))
    r.add_argument('--workdir', default=None)
    r.set_defaults(func=cmd_run)

    args = ap.parse_args()
    try:
        out = args.func(args)
    except Exception as exc:
        import traceback
        out = {'ok': False, 'error': f'{type(exc).__name__}: {exc}',
               'traceback': traceback.format_exc()[-1500:]}
    print(json.dumps(out, ensure_ascii=False, indent=1, default=str))
    raise SystemExit(0 if out.get('ok') else 1)


if __name__ == '__main__':
    main()
