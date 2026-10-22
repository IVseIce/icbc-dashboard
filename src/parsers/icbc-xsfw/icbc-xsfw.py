#!/home/isee/.venvs/data-tools/bin/python
# -*- coding: utf-8 -*-
"""icbc-xsfw —— 阜阳分行「网点线上服务能力」提升通报：取信 → 计算 → 出表 → 出文案 → 回信。

脚本只做确定性的事：IMAP 取信/下载附件、按内容识别源表、按固定口径计算、
按固定版式写 xlsx、按固定模板生成文案、SMTP 回信。判断类的事（选哪封邮件、
文案措辞微调）交给模型，见同目录 SKILL.md。

用法
----
  icbc-xsfw.py fetch  --uid U [--uid U ...] [--root DIR] [--workdir DIR]
  icbc-xsfw.py build  --workdir DIR [--source DIR] [--no-text]
  icbc-xsfw.py send   --workdir DIR [--to ADDR] [--dry-run]
  icbc-xsfw.py run    --uid U [--uid U ...] [--root DIR]      # fetch + build，不发信

输出一律 JSON（stdout），便于模型直接读取。任何源表缺失/不匹配都返回 ok=false 并停止。
"""
from __future__ import annotations

import argparse
import csv as csvmod
import email
import imaplib
import io
import json
import os
import re
import smtplib
import socket
import ssl
import sys
import warnings
from datetime import date, datetime, timedelta
from email.header import make_header, decode_header
from email.message import EmailMessage
from email.utils import formatdate, parsedate_to_datetime, parseaddr
from pathlib import Path
from zoneinfo import ZoneInfo

try:
    import openpyxl
    import yaml
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    print(json.dumps({'ok': False, 'error': f'缺少依赖: {exc}；请用 /home/isee/.venvs/data-tools/bin/python 运行'},
                     ensure_ascii=False))
    raise SystemExit(2)

warnings.filterwarnings('ignore')

TZ = ZoneInfo('Asia/Shanghai')
DSH_HOME = Path(os.environ.get('DSH_HOME') or (Path.home() / '.dsh'))
SKILL_DIR = Path(__file__).resolve().parent
FALLBACK_TEMPLATE = SKILL_DIR / 'assets' / 'template.xlsx'
DEFAULT_ROOT = Path('/mnt/d/资料/ICBC/automations')
WORKDIR_SUFFIX = '_网点线上服务能力提升通报'
REPORT_STEM = '阜阳分行“网点线上服务能力”提升通报'

# ---------------------------------------------------------------- 口径常量
BRANCH_MAP = {
    '颍州': ['颍州支行营业室', '金鹰支行', '南站支行', '清河支行', '师范大学支行', '城建支行', '文峰支行'],
    '颍泉': ['颍泉支行营业室', '临沂商城支行', '瑶海支行', '人民路支行', '大楼支行', '小隅首支行', '牡丹支行'],
    '颍东': ['颍河支行', '颍东支行营业室', '青年支行', '烟厂支行', '岳家湖支行', '汇通支行'],
    '界首': ['界首支行营业室', '界首城建支行', '界首科创支行', '界首新街口支行'],
    '临泉': ['临泉支行营业室', '临泉文明支行', '临泉政务新区支行', '临泉姜尚支行'],
    '阜南': ['阜南支行营业室', '阜南百货大楼支行', '阜南北城支行', '阜南胜利路支行'],
    '颍上': ['颍上支行营业室', '颍上北关支行', '颍上十字街支行', '颍上管仲支行'],
    '太和': ['太和支行营业室', '太和人民路支行', '太和大桥支行', '太和华源支行'],
}
# 通报模板行名 → 源表网点简称（网点更名/历史命名差异）。
# 模板是手工维护的，可能仍带着旧名，所以入模板的旧名也要认。
TEMPLATE_ALIAS = {
    '阜阳颍东颍上路支行': '岳家湖支行',
    '阜阳界首影剧院支行': '界首科创支行',
}
# 输出时把模板里的旧行名统一改成现名（A 列与 N 列镜像都要改）。
ROW_RENAME = {
    '阜阳颍东颍上路支行': '阜阳岳家湖支行',
    '阜阳界首影剧院支行': '阜阳界首科创支行',
}
MIN_CUSTOMERS = 5          # 当日到店客户数 < 此值 → 轮休，不算动户率
LUNCH_TEXT = '轮休'
# 法定节假日（当前不作为硬性条件，仅按「当日到店 < MIN_CUSTOMERS」判轮休）。留作后续启用。
HOLIDAYS: set[str] = set()

FONT_NAME, FONT_SIZE = '微软雅黑', 11
COLOR_TOP, COLOR_BOTTOM = 'FFFF0000', 'FF008000'
PCT_COLS = {'B', 'C', 'E', 'J', 'M'}
NUM_COLS = {'D', 'F', 'G', 'H', 'I', 'K', 'L'}
ALL_COLS = ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M']
FIRST_DATA_ROW, LAST_DATA_ROW, TOTAL_ROW = 4, 43, 44

PROVIDERS = {
    '163': ('imap.163.com', 993, 'smtp.163.com', 465),
    '126': ('imap.126.com', 993, 'smtp.126.com', 465),
    'qq': ('imap.qq.com', 993, 'smtp.qq.com', 465),
    'sina': ('imap.sina.com', 993, 'smtp.sina.com', 465),
    'aliyun': ('imap.aliyun.com', 993, 'smtp.aliyun.com', 465),
    'gmail': ('imap.gmail.com', 993, 'smtp.gmail.com', 465),
    'outlook': ('outlook.office365.com', 993, 'smtp.office365.com', 587),
    'icloud': ('imap.mail.me.com', 993, 'smtp.mail.me.com', 587),
}


# ================================================================ 邮箱账号
def _read_row_config() -> dict:
    out: dict = {}
    for cfg in sorted((DSH_HOME / 'profiles').glob('*/cordis.patch.yml')):
        try:
            data = yaml.safe_load(cfg.read_text('utf-8')) or []
        except Exception:
            continue
        if not isinstance(data, list):
            continue
        for entry in data:
            if isinstance(entry, dict) and entry.get('id') == 'tool-email':
                conf = entry.get('config')
                if isinstance(conf, dict):
                    out.update(conf)
    return out


def _read_page_config() -> dict:
    try:
        data = yaml.safe_load((DSH_HOME / 'settings.yaml').read_text('utf-8')) or {}
    except Exception:
        return {}
    ns = data.get('dsh-email')
    return ns if isinstance(ns, dict) else {}


def _merge_config(row: dict, page: dict) -> dict:
    cfg = dict(row)
    for key in ('provider', 'user', 'password', 'inboxFolder', 'downloadDir'):
        value = page.get(key)
        if isinstance(value, str) and value.strip():
            cfg[key] = value.strip()
    return cfg


def _pick_account(cfg: dict) -> dict:
    entries, default = None, None
    raw_yaml = cfg.get('accountsYaml') if isinstance(cfg.get('accountsYaml'), str) else ''
    if (raw_yaml or '').strip():
        doc = yaml.safe_load(raw_yaml) or {}
        if isinstance(doc, dict):
            doc = dict(doc)
            default = doc.pop('defaultAccount', None)
            entries = doc
    elif isinstance(cfg.get('accounts'), dict) and cfg['accounts']:
        entries = cfg['accounts']
    if not entries:
        return cfg
    name = cfg.get('defaultAccount') or default or ('default' if 'default' in entries else sorted(entries)[0])
    account = entries.get(name)
    if not isinstance(account, dict):
        raise RuntimeError(f'accountsYaml 里没有默认账号 "{name}"')
    merged = {k: v for k, v in account.items() if v not in (None, '')}
    for key in ('provider', 'user', 'password', 'inboxFolder', 'imap', 'smtp'):
        if key not in merged and cfg.get(key):
            merged[key] = cfg[key]
    return merged


def _env_password() -> tuple[str, str] | None:
    pw = os.environ.get('DSH_EMAIL_PASSWORD')
    if pw and pw.strip():
        return pw.strip(), 'DSH_EMAIL_PASSWORD 环境变量'
    for pid in os.listdir('/proc'):
        if not pid.isdigit():
            continue
        try:
            cmd = Path(f'/proc/{pid}/cmdline').read_bytes().replace(b'\0', b' ').decode('utf-8', 'ignore')
            if 'dsh' not in cmd or 'web' not in cmd:
                continue
            for kv in Path(f'/proc/{pid}/environ').read_bytes().split(b'\0'):
                if kv.startswith(b'DSH_EMAIL_PASSWORD='):
                    return kv.split(b'=', 1)[1].decode('utf-8', 'ignore').strip(), '运行中 dsh web 的环境'
        except (OSError, PermissionError):
            continue
    return None


def resolve_account() -> dict:
    cfg = _pick_account(_merge_config(_read_row_config(), _read_page_config()))
    provider = (cfg.get('provider') or '163').strip() or '163'
    imap_host, imap_port, smtp_host, smtp_port = PROVIDERS.get(provider, PROVIDERS['163'])
    imap_conf = cfg.get('imap') if isinstance(cfg.get('imap'), dict) else {}
    smtp_conf = cfg.get('smtp') if isinstance(cfg.get('smtp'), dict) else {}
    pwd = cfg.get('password')
    if isinstance(pwd, str) and pwd.strip():
        password, source = pwd.strip(), '插件配置（设置页/行配置）'
    else:
        found = _env_password()
        if not found:
            raise RuntimeError('取不到邮箱授权码：请在邮件设置页填写，或设置 DSH_EMAIL_PASSWORD')
        password, source = found
    return {
        'provider': provider,
        'user': str(cfg.get('user') or '').strip(),
        'imap_host': str(imap_conf.get('host') or '').strip() or imap_host,
        'imap_port': int(imap_conf.get('port') or imap_port),
        'smtp_host': str(smtp_conf.get('host') or '').strip() or smtp_host,
        'smtp_port': int(smtp_conf.get('port') or smtp_port),
        'password': password,
        'password_source': source,
        'inbox': str(cfg.get('inboxFolder') or 'INBOX'),
    }


def force_ipv4() -> None:
    original = socket.getaddrinfo

    def ipv4_only(host, port, family=0, type=0, proto=0, flags=0):
        return original(host, port, socket.AF_INET, type, proto, flags)

    socket.getaddrinfo = ipv4_only


def connect(acct: dict) -> imaplib.IMAP4_SSL:
    force_ipv4()
    imap = imaplib.IMAP4_SSL(acct['imap_host'], acct['imap_port'],
                             ssl_context=ssl.create_default_context(), timeout=60)
    imap.login(acct['user'], acct['password'])
    try:
        imap.xatom('ID', '("name" "dsh-icbc-xsfw" "version" "1.0" "vendor" "dsh")')
    except Exception:
        pass
    return imap


def _decode(value) -> str:
    if value is None:
        return ''
    try:
        return str(make_header(decode_header(str(value))))
    except Exception:
        return str(value)


# ================================================================ 取信/下载
def fetch_message(imap, uid: str) -> dict:
    typ, body = imap.uid('FETCH', uid, '(BODY.PEEK[])')
    if typ != 'OK' or not body or not isinstance(body[0], tuple):
        raise RuntimeError(f'UID {uid} 取信失败')
    msg = email.message_from_bytes(body[0][1])
    try:
        received = parsedate_to_datetime(msg['Date']).astimezone(TZ)
    except Exception:
        received = None
    atts = []
    for part in msg.walk():
        filename = _decode(part.get_filename())
        if not filename:
            continue
        payload = part.get_payload(decode=True)
        if payload is None:
            continue
        atts.append({'filename': filename, 'blob': payload})
    return {
        'uid': uid,
        'subject': _decode(msg['Subject']),
        'from': _decode(msg['From']),
        'message_id': (msg['Message-ID'] or '').strip(),
        'received': received,
        'received_date': received.date().isoformat() if received else None,
        'attachments': atts,
    }


_MONTHS = ('Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec')


def _imap_date(d: date) -> str:
    """IMAP 的日期格式 28-Sep-2026（英文月，不能用 strftime 以免受 locale 影响）。"""
    return f'{d.day:02d}-{_MONTHS[d.month - 1]}-{d.year}'


def list_recent(imap, since: date, cap: int = 400) -> list[dict]:
    """列出 SINCE 之后的邮件抬头（主题/发件人/日期），不下载正文与附件。

    必须用 UID SEARCH：普通 SEARCH 回的是序号，拿去做 UID FETCH 会取不到东西。
    """
    typ, data = imap.uid('SEARCH', None, f'(SINCE {_imap_date(since)})')
    if typ != 'OK':
        raise RuntimeError(f'IMAP SEARCH 失败：{typ}')
    uids = (data[0] or b'').split()[-cap:]
    out = []
    for u in uids:
        uid = u.decode()
        try:
            typ, d = imap.uid('FETCH', uid, '(BODY.PEEK[HEADER.FIELDS (SUBJECT DATE FROM)])')
        except Exception:
            continue
        if typ != 'OK' or not d or not isinstance(d[0], tuple):
            continue
        msg = email.message_from_bytes(d[0][1])
        try:
            received = parsedate_to_datetime(msg['Date']).astimezone(TZ)
        except Exception:
            received = None
        out.append({'uid': uid, 'subject': _decode(msg['Subject']), 'from': _decode(msg['From']),
                    'received': received,
                    'received_date': received.date().isoformat() if received else None})
    out.sort(key=lambda x: x['received'] or datetime.min.replace(tzinfo=TZ))
    return out


def pool_attachments(imap, uids: list[str], cache: dict) -> list[dict]:
    """把若干封邮件的附件汇总成一个池子（同 uid 只取一次）。"""
    out = []
    for uid in uids:
        if uid not in cache:
            cache[uid] = fetch_message(imap, uid)
        out.extend(cache[uid]['attachments'])
    return out


def pool_check(imap, uids: list[str], cache: dict) -> dict:
    """按内容核验一个附件池能否凑齐所有必需角色。"""
    try:
        cls = classify(pool_attachments(imap, uids, cache))
    except Exception as exc:
        return {'complete': False, 'problems': [f'{type(exc).__name__}: {exc}'],
                'notices': [], 'attachment_count': 0, 'roles': {}}
    cls['complete'] = all(r in cls['roles'] for r in REQUIRED_ROLES)
    return cls


def strip_blobs(cls: dict) -> dict:
    """输出用：把附件字节去掉，只留文件名与日期（否则 JSON 会被二进制撑爆）。"""
    return {
        'complete': bool(cls.get('complete')),
        'problems': cls.get('problems', []),
        'notices': cls.get('notices', []),
        'unknown_files': cls.get('unknown_files', []),
        'attachment_count': cls.get('attachment_count', 0),
        'roles': {k: {'file': v.get('filename'), 'data_date': v.get('data_date')}
                  for k, v in cls.get('roles', {}).items()},
    }


def minimal_pick(imap, uids: list[str], cache: dict) -> tuple[list[str], set, dict]:
    """从一批邮件里挑出**最少**的几封凑齐必需角色。

    不能按单封邮件判断角色：「当月基准 CSV」只有把几份 CSV 放在一起比日期才认得出来。
    所以先把整批附件汇总做一次 classify，再看每个入选角色是谁供的，只保留那些邮件。
    """
    pairs: list[tuple[str, dict]] = []
    for u in uids:
        if u not in cache:
            cache[u] = fetch_message(imap, u)
        for a in cache[u]['attachments']:
            pairs.append((u, a))
    try:
        cls = classify([a for _, a in pairs])
    except Exception:
        return [], set(), {'roles': {}, 'problems': []}
    owner: dict[str, str] = {}
    for u, a in pairs:
        owner.setdefault(a['filename'], u)
    picked: list[str] = []
    for role in REQUIRED_ROLES + ('template',):
        entry = cls['roles'].get(role)
        if not entry:
            continue
        u = owner.get(entry['filename'])
        if u and u not in picked:
            picked.append(u)
    ordmap = {u: i for i, u in enumerate(uids)}
    picked.sort(key=lambda u: ordmap.get(u, 0))
    return picked, set(cls['roles']), cls


# ================================================================ 源表识别
def _csv_rows(blob: bytes) -> list[list[str]]:
    for enc in ('utf-8-sig', 'utf-8', 'gbk', 'gb18030'):
        try:
            text = blob.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise RuntimeError('CSV 编码无法识别')
    return [row for row in csvmod.reader(io.StringIO(text))]


def probe_csv(blob: bytes) -> dict | None:
    """客户经理明细 CSV → {kind, data_date, rows}；不是这张表返回 None。"""
    try:
        raw = _csv_rows(blob)
    except Exception:
        return None
    if len(raw) < 2:
        return None
    m = re.search(r'(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日', ','.join(raw[0]))
    header = [c.strip() for c in raw[1]]
    if not m or '客户总数' not in header or '合计认证好友数(去重)' not in header:
        return None
    idx = {c: i for i, c in enumerate(header)}
    need = ['机构名称', '客户总数', '合计认证好友数(去重)']
    if not all(k in idx for k in need):
        return None
    agg: dict[str, dict[str, float]] = {}
    for row in raw[2:]:
        if len(row) <= max(idx[k] for k in need):
            continue
        org = row[idx['机构名称']].strip()
        if not org:
            continue
        def num(key):
            try:
                return float(str(row[idx[key]]).strip() or 0)
            except ValueError:
                return 0.0
        a = agg.setdefault(org, {'khs': 0.0, 'rz': 0.0})
        a['khs'] += num('客户总数')
        a['rz'] += num('合计认证好友数(去重)')
    return {'kind': 'staff_csv', 'data_date': f'{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}',
            'nets': agg}


def probe_donghu(blob: bytes) -> dict | None:
    """到店客群手机银行渗透率和动户率表 → {kind, data_date, nets, total}。"""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)
    except Exception:
        return None
    ws = wb.worksheets[0]
    title = str(ws.cell(row=1, column=1).value or '')
    if '到店客群' not in title or '动户率' not in title:
        return None
    m = re.search(r'日期[:：]\s*(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日',
                  str(ws.cell(row=2, column=4).value or ''))
    nets: dict[str, dict[str, float]] = {}
    total: dict[str, float] = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        name = row[0]
        if not name:
            continue
        vals = {'B': float(row[1] or 0), 'act': float(row[3] or 0), 'I': float(row[8] or 0)}
        if str(name).strip() == '合计':
            total = vals
        else:
            nets[str(name).strip()] = vals
    if not total:
        return None
    return {'kind': 'donghu',
            'data_date': f'{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}' if m else None,
            'nets': nets, 'total': total}


def probe_auth(blob: bytes) -> dict | None:
    """营销授权通报 → {kind, data_date, nets}。表头是 zhihang/wangdian/khs/qys/ghs/ghqys..."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True)
    except Exception:
        return None
    ws = wb.worksheets[0]
    head = [str(c.value or '').strip() for c in ws[1]]
    need = ['wangdian', 'khs', 'qys', 'ghs', 'ghqys']
    if not all(k in head for k in need):
        return None
    idx = {k: head.index(k) for k in need}
    dt = head.index('DATA_DT') if 'DATA_DT' in head else None
    nets: dict[str, dict[str, float]] = {}
    data_date = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[idx['wangdian']]:
            continue
        name = str(row[idx['wangdian']]).strip()
        def num(key, default=0.0):
            try:
                return float(row[idx[key]] or default)
            except (TypeError, ValueError):
                return default
        nets[name] = {'khs': num('khs'), 'qys': num('qys'), 'ghs': num('ghs'), 'ghqys': num('ghqys')}
        if dt is not None and not data_date and row[dt]:
            data_date = str(row[dt])[:10]
    return {'kind': 'auth', 'data_date': data_date, 'nets': nets}


def probe_template(blob: bytes) -> dict | None:
    """汇总通报表样式（空白模板）→ {kind, title, labels}。"""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(blob))
    except Exception:
        return None
    ws = wb.worksheets[0]
    title = str(ws.cell(row=1, column=1).value or '')
    if '网点线上服务能力' not in title or '提升通报' not in title:
        return None
    labels = [str(ws.cell(row=r, column=1).value or '').strip() for r in range(FIRST_DATA_ROW, TOTAL_ROW)]
    # 模板必须是空的；已经填过数的（例如别人发来的某日成表）不算模板，免得张冠李戴
    filled = any(ws.cell(row=r, column=c).value not in (None, '')
                 for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1) for c in range(2, 14))
    return {'kind': 'filled_report' if filled else 'template',
            'title': title, 'labels': [x for x in labels if x]}


PROBES = (probe_donghu, probe_auth, probe_template, probe_csv)

# 缺任何一个都算不出完整通报表（模板可选，见 ensure_template）
REQUIRED_ROLES = ('donghu_day', 'donghu_prev', 'staff_cur', 'staff_prev', 'staff_month', 'auth')
KEYWORD = '线上服务'

# 回信固定收件人：原发件人 + 这三个人；发件人本身在名单里就只发一遍
FIXED_RECIPIENTS = (
    'zhaili.fy@ah.icbc.com.cn',
    'zhangxiang.fy@ah.icbc.com.cn',
    'yaohongcheng.fy@ah.icbc.com.cn',
)


def dedup_addresses(addrs) -> list[str]:
    """去重同时保持顺序；邮箱大小写不敏感。"""
    seen, out = set(), []
    for a in addrs:
        a = (a or '').strip()
        if not a or a.lower() in seen:
            continue
        seen.add(a.lower())
        out.append(a)
    return out


def classify(attachments: list[dict]) -> dict:
    """把附件按内容分成各角色；任何缺失/歧义都在 problems 里说明，不猜。"""
    found: dict[str, list[dict]] = {}
    unknown: list[str] = []
    for att in attachments:
        got = None
        for probe in PROBES:
            got = probe(att['blob'])
            if got:
                break
        if not got:
            if att['filename'].lower().endswith(('.xls', '.xlsx', '.xlsm', '.csv')):
                unknown.append(att['filename'])
            continue
        got['filename'] = att['filename']
        got['blob'] = att['blob']
        found.setdefault(got['kind'], []).append(got)

    roles: dict[str, dict] = {}
    problems: list[str] = []
    notices: list[str] = []

    # 认不出的表格类附件：可能是源表换了版式，必须让用户看到，不能默默丢掉
    if unknown:
        problems.append('这些表格附件认不出是哪张源表，请确认是否换了版式或漏了对应关系：'
                        + '、'.join(unknown))
    filled = [f['filename'] for f in found.get('filled_report', [])]
    if filled:
        notices.append('这些是已填过数的通报表（非空白模板），已忽略：' + '、'.join(filled))

    donghu = found.get('donghu', [])
    if len(donghu) < 2:
        problems.append(f'到店客群表需要 2 个（当日+前一日），实际 {len(donghu)} 个')
    else:
        donghu.sort(key=lambda x: x['data_date'] or '')
        roles['donghu_day'], roles['donghu_prev'] = donghu[-1], donghu[-2]
        if len(donghu) > 2:
            problems.append(f'到店客群表多于 2 个，已取日期最新的两个：'
                            f'{[d["filename"] for d in donghu]}')

    csvs = sorted(found.get('staff_csv', []), key=lambda x: x['data_date'] or '')
    if not csvs:
        problems.append('没有识别到「客户经理明细」CSV（需含 客户总数 / 合计认证好友数(去重) 两列）')
    else:
        roles['staff_cur'] = csvs[-1]
        cur_dt = date.fromisoformat(roles['staff_cur']['data_date'])
        month_base = [c for c in csvs[:-1]
                      if date.fromisoformat(c['data_date']).year * 12 + date.fromisoformat(c['data_date']).month
                      < cur_dt.year * 12 + cur_dt.month]
        if month_base:
            roles['staff_month'] = month_base[-1]
            rest = [c for c in csvs[:-1] if c is not roles['staff_month']]
        else:
            rest = csvs[:-1]
            problems.append('没有跨月的明细 CSV，无法算「当月添客」（G 列）')
        if rest:
            roles['staff_prev'] = rest[-1]
        else:
            problems.append('没有可用的「上一日」明细 CSV，无法算「当日添客」（F 列）')

    auth = found.get('auth', [])
    if len(auth) != 1:
        problems.append(f'营销授权通报需要 1 个，实际 {len(auth)} 个')
    else:
        roles['auth'] = auth[0]

    tpl = found.get('template', [])
    if len(tpl) > 1:
        problems.append(f'空白通报表模板多于 1 个，无法确定用哪个：{[t["filename"] for t in tpl]}')
    elif len(tpl) == 1:
        roles['template'] = tpl[0]
    else:
        # 模板不是必须的：邮件没带就用技能内置的（标题日期按业务日重写）
        notices.append('邮件里没有空白通报表模板，将使用技能内置模板 assets/template.xlsx')

    for r in ('staff_cur', 'staff_prev', 'staff_month'):
        v = roles.get(r)
        if v:
            v['nets'] = {k.replace('阜阳', ''): x for k, x in v['nets'].items()}

    return {'ok': not problems, 'roles': roles, 'unknown_files': unknown,
            'problems': problems, 'notices': notices,
            'attachment_count': len(attachments)}


# ================================================================ 网点匹配
def norm(name: str) -> str:
    s = str(name).strip().replace('阜阳', '')
    if s.endswith('营业室') and len(s) > 3:
        s = s[:-3]
    return s.replace('师范学院', '师范大学')


def net_keys(branch: str, net: str) -> list[str]:
    keys = [norm(net)]
    if not net.startswith(branch):
        keys.append(norm(branch + net))
    return keys


def lookup_for(raw_names) -> dict[str, str]:
    have = {norm(n) for n in raw_names}
    primary, fallback = {}, {}
    for branch, nets in BRANCH_MAP.items():
        for net in nets:
            nk = norm(net)
            full = norm(branch + net) if not net.startswith(branch) else nk
            primary.setdefault(full, net)
            fallback.setdefault(nk, net)
    primary = {k: v for k, v in primary.items() if k in have}
    fallback = {k: v for k, v in fallback.items() if k in have}
    return {**fallback, **primary}


def reindex(raw: dict, dest: dict[str, str]) -> dict[str, dict]:
    out: dict[str, dict] = {}
    for src, vals in raw.items():
        net = dest.get(norm(src))
        if net:
            out.setdefault(net, {}).update(vals)
    return out


def short_of(label: str, net: str) -> str:
    """文案里的网点简称：X支行营业室 → X营业室；其余去掉「支行」后缀。

    是否带「营业室」以通报模板行名为准（BRANCH_MAP 里的汇通/文峰不带营业室）。
    名称用源表口径（如 岳家湖/界首科创），模板里的旧名只在 xlsx 行名上保留。
    """
    if net == '师范大学支行':          # 历史命名差异
        return '师范学院'
    lab = str(label).replace('阜阳', '')
    has_ys = net.endswith('支行营业室') or lab.endswith('支行营业室')
    if net.endswith('支行营业室'):
        base = net[:-5]
    elif net.endswith('支行'):
        base = net[:-2]
    else:
        base = net
    return base + '营业室' if has_ys else base


# ================================================================ 计算
def compute(roles: dict) -> dict:
    day, prev = roles['donghu_day'], roles['donghu_prev']
    day_n = {k: v for k, v in day['nets'].items()}
    prev_n = {k: v for k, v in prev['nets'].items()}
    day_m = reindex(day_n, lookup_for(day_n))
    prev_m = reindex(prev_n, lookup_for(prev_n))
    cur = reindex(roles['staff_cur']['nets'], lookup_for(roles['staff_cur']['nets']))
    pre = reindex(roles['staff_prev']['nets'], lookup_for(roles['staff_prev']['nets']))
    mon = reindex(roles['staff_month']['nets'], lookup_for(roles['staff_month']['nets']))
    auth_all = roles['auth']['nets']
    auth = reindex(auth_all, lookup_for(auth_all))

    # 模板行名先统一成现名（旧名照旧能匹配，见 TEMPLATE_ALIAS）
    labels = [ROW_RENAME.get(x, x) for x in roles['template']['labels']]
    tpl_to_net: dict[str, str] = {}
    nets_all = [n for ns in BRANCH_MAP.values() for n in ns]
    for label in labels:
        if label in TEMPLATE_ALIAS:
            tpl_to_net[label] = TEMPLATE_ALIAS[label]
            continue
        hit = next((n for n in nets_all
                    if norm(label) in net_keys(next(b for b, ns in BRANCH_MAP.items() if n in ns), n)), None)
        if hit is None:
            raise RuntimeError(f'通报模板行无法匹配网点：{label}')
        tpl_to_net[label] = hit

    rows, warns = [], []
    for label in labels:
        net = tpl_to_net[label]
        branch = next(b for b, ns in BRANCH_MAP.items() if net in ns)
        rec: dict = {'label': label, 'net': net, 'branch': branch, 'short': short_of(label, net)}
        d, p = day_m.get(net), prev_m.get(net)
        if d and p:
            cust = d['B'] - p['B']
            act = d['I'] * d['B'] - p['I'] * p['B']
            rec['C'] = d['I']
            rec['cust_delta'] = cust
            if cust < MIN_CUSTOMERS:
                rec['B'] = None
                warns.append(f'{label}: 当日到店客户 {cust:g} 人（<{MIN_CUSTOMERS}），按轮休处理')
            else:
                rec['B'] = act / cust
        else:
            rec['B'] = rec['C'] = None
            warns.append(f'{label}: 到店客群表缺行（当日{"有" if d else "无"}/前一日{"有" if p else "无"}）')
        c = cur.get(net)
        if c:
            rec['D'], rec['RZ'] = c['khs'], c['rz']
            rec['E'] = (c['rz'] / c['khs']) if c['khs'] else None
            rec['F'] = c['khs'] - pre.get(net, {}).get('khs', c['khs'])
            rec['G'] = c['khs'] - mon.get(net, {}).get('khs', c['khs'])
        else:
            rec['D'] = rec['E'] = rec['F'] = rec['G'] = None
            warns.append(f'{label}: 客户经理明细缺行')
        a = auth.get(net)
        if a:
            rec['H'], rec['I'] = a['khs'], a['qys']
            rec['K'], rec['L'] = a['ghs'], a['ghqys']
            rec['J'] = (a['qys'] / a['khs']) if a['khs'] else None
            rec['M'] = (a['ghqys'] / a['ghs']) if a['ghs'] else None
        else:
            for k in 'HIJKLM':
                rec[k] = None
            warns.append(f'{label}: 营销授权通报缺行')
        rows.append(rec)

    cust = day['total']['B'] - prev['total']['B']
    act = day['total']['I'] * day['total']['B'] - prev['total']['I'] * prev['total']['B']
    tot: dict = {'label': '合计', 'net': '合计', 'branch': '', 'short': '合计'}
    tot['B'] = (act / cust) if cust else None
    tot['C'] = day['total']['I'] or None
    tot['cust_delta'] = cust
    for k in 'DFGHIKL':
        vals = [r[k] for r in rows if r.get(k) is not None]
        tot[k] = sum(vals) if vals else None
    tot_rz = sum(r['RZ'] for r in rows if r.get('RZ') is not None)
    tot['E'] = (tot_rz / tot['D']) if tot.get('D') else None
    tot['J'] = (tot['I'] / tot['H']) if tot.get('H') else None
    tot['M'] = (tot['L'] / tot['K']) if tot.get('K') else None
    rows.append(tot)

    return {'day_date': day['data_date'], 'prev_date': prev['data_date'],
            'staff_dates': {'cur': roles['staff_cur']['data_date'],
                            'prev': roles['staff_prev']['data_date'],
                            'month': roles['staff_month']['data_date']},
            'auth_date': roles['auth']['data_date'],
            'rows': rows, 'warnings': warns}


# ================================================================ 写表
def write_report(template_blob: bytes, data: dict, out_path: Path) -> None:
    wb = openpyxl.load_workbook(io.BytesIO(template_blob))
    ws = wb.worksheets[0]
    for i, rec in enumerate(data['rows']):
        r = FIRST_DATA_ROW + i
        # A 列行名写成现名，N 列是镜像列，跟着一起改
        ws[f'A{r}'].value = rec['label']
        if ws[f'N{r}'].value is not None:
            ws[f'N{r}'].value = rec['label']
        for col in ALL_COLS:
            val = rec.get(col)
            cell = ws[f'{col}{r}']
            if val is None:
                cell.value = LUNCH_TEXT if col == 'B' else None
            else:
                cell.value = round(val, 6)
            cell.number_format = '0.00%' if col in PCT_COLS else '#,##0'
        if rec.get('B') is None:
            ws[f'B{r}'].number_format = 'General'
            ws[f'B{r}'].alignment = Alignment(horizontal='center')

    # 字体：全表微软雅黑 11（合并区里被合并掉的格子按坐标补）
    def setfont(cell):
        f = cell.font
        cell.font = Font(name=FONT_NAME, sz=FONT_SIZE, b=f.b, i=f.i, u=f.u, color=f.color)

    for row in ws.iter_rows(min_row=1, max_row=TOTAL_ROW, min_col=1, max_col=14):
        for cell in row:
            setfont(cell)
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                try:
                    setfont(ws.cell(row=r, column=c))
                except AttributeError:
                    pass

    # 排名：每列最大 5 标红、最小 5 标绿；不足 10 个数值时红色优先（红5绿4）
    for col in ALL_COLS:
        pairs = [(r, ws[f'{col}{r}'].value) for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1)]
        nums = [(r, v) for r, v in pairs if isinstance(v, (int, float))]
        for r, _ in pairs:
            f = ws[f'{col}{r}'].font
            ws[f'{col}{r}'].font = Font(name=FONT_NAME, sz=FONT_SIZE, b=f.b, i=f.i, u=f.u)
        if len(nums) < 2:
            continue
        nums.sort(key=lambda p: p[1])
        if len(nums) >= 10:
            hi = lo = 5
        else:
            hi = min(5, len(nums))
            lo = max(0, len(nums) - hi)
        for r, _ in nums[len(nums) - hi:]:
            ws[f'{col}{r}'].font = Font(name=FONT_NAME, sz=FONT_SIZE, bold=True, color=COLOR_TOP)
        for r, _ in nums[:lo]:
            ws[f'{col}{r}'].font = Font(name=FONT_NAME, sz=FONT_SIZE, bold=True, color=COLOR_BOTTOM)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


# ================================================================ 文案
def trunc2(v: float) -> float:
    import math
    return math.floor(v * 100 + 1e-9) / 100


def build_text(data: dict) -> dict:
    """按 2026-10-13 用户确认的新版式生成文案：
       1 段 总体（当月动户率 / 当日企微添客 / 营销授权率）
       2 段 到店动户率 前3后2
       3 段 企微添客（当日添客）前3后3
       4 段 营销授权（签约率垫底2 + 管户签约率靠后2 + 支行观察）"""
    rows = [r for r in data['rows'] if r['net'] != '合计']
    tot = data['rows'][-1]
    notes: list[str] = []

    def pct(v):
        return f'{v * 100:.2f}'

    d = date.fromisoformat(data['day_date'])
    lines = []

    # —— 1 总体
    lines.append(f'{d.month}月{d.day}日，全行当月到店客户手机银行动户率{pct(tot["C"])}%，'
                 f'当日企微添客{tot["F"]:g}户，全量客户营销授权率{pct(tot["J"])}%。')

    # —— 2 到店动户率（样本太少退回当月动户率）
    bvals = sorted([r for r in rows if isinstance(r['B'], (int, float))], key=lambda r: -r['B'])
    cvals = sorted([r for r in rows if isinstance(r['C'], (int, float))], key=lambda r: -r['C'])
    if len(bvals) >= 5:
        use, label, key = bvals, '到店动户率', 'B'
    else:
        use, label, key = cvals, '当月动户率', 'C'
        notes.append(f'当日动户率只有 {len(bvals)} 家有值（其余轮休），该段改用当月动户率')
    if len(use) >= 3:
        lines.append(
            f'{label}方面，{use[0]["short"]}{pct(use[0][key])}%领跑[强][强]，'
            f'{use[1]["short"]}{pct(use[1][key])}%、{use[2]["short"]}{pct(use[2][key])}%表现靠前；'
            f'{use[-1]["short"]}{pct(use[-1][key])}%、{use[-2]["short"]}{pct(use[-2][key])}%排名靠后[加油][加油]')

    # —— 3 企微添客（当日添客 F 列，前3 / 后3）
    fdesc = sorted([r for r in rows if isinstance(r['F'], (int, float))], key=lambda r: -r['F'])
    fasc = sorted([r for r in rows if isinstance(r['F'], (int, float))], key=lambda r: r['F'])
    if len(fdesc) >= 3:
        top, bot = fdesc[:3], fasc[:3]          # 领先取最高3；落后取最低3（并列时按模板行序）
        seg = ('企微添客方面，'
               + '、'.join(f'{r["short"]}{r["F"]:g}' for r in top) + '当日领先全行。')
        if bot:
            seg += '、'.join(f'{r["short"]}{r["F"]:g}' for r in bot) + '暂时落后[心碎][心碎]'
        lines.append(seg)

    # —— 4 营销授权
    jvals = sorted([r for r in rows if isinstance(r['J'], (int, float))], key=lambda r: r['J'])
    mvals = sorted([r for r in rows if isinstance(r['M'], (int, float))], key=lambda r: r['M'])
    if len(jvals) >= 2 and len(mvals) >= 2:
        seg = (f'营销授权方面，'
               f'{jvals[0]["short"]}{pct(jvals[0]["J"])}%、{jvals[1]["short"]}{pct(jvals[1]["J"])}%签约率垫底；'
               f'{mvals[0]["short"]}{pct(mvals[0]["M"])}%、{mvals[1]["short"]}{pct(mvals[1]["M"])}%管户签约率靠后。')
        agg: dict[str, dict[str, float]] = {}
        for r in rows:
            a = agg.setdefault(r['branch'], {'G': 0.0})
            if isinstance(r['G'], (int, float)):
                a['G'] += r['G']
            for k in 'IHKL':
                a.setdefault(k, 0.0)
                if isinstance(r[k], (int, float)):
                    a[k] += r[k]
        if agg:
            top_add = max(agg, key=lambda b: agg[b]['G'])
            jr = sorted(agg, key=lambda b: agg[b]['I'] / agg[b]['H'] if agg[b]['H'] else 9)
            mr = sorted(agg, key=lambda b: agg[b]['L'] / agg[b]['K'] if agg[b]['K'] else 9)
            if top_add in jr[:2] and top_add in mr[:2]:
                seg += f'{top_add}支行添客规模第一但两项签约率均偏低，需抓转化[加油][加油]'
            elif top_add in jr[:2]:
                seg += f'{top_add}支行添客规模第一但签约率偏低，需抓转化[加油][加油]'
            else:
                seg += '请各支行持续抓好存量客户签约转化[加油][加油]'
        lines.append(seg)

    text = '\n\n'.join(lines) + '\n'
    return {'text': text, 'notes': notes}


# ================================================================ 子命令
def _workdir_for(root: Path, received: date) -> Path:
    return root / f'{received.strftime("%Y%m%d")}{WORKDIR_SUFFIX}'


def _save(path: Path, blob: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(blob)


def _retitle(blob: bytes, biz_date: str) -> tuple[bytes, str]:
    """把模板 A1 标题里的（MM.DD）改成业务日；没有就补一个。返回 (新字节, 新标题)。"""
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    ws = wb.worksheets[0]
    old = str(ws.cell(row=1, column=1).value or REPORT_STEM)
    d = date.fromisoformat(biz_date)
    stamp = f'（{d.month:02d}.{d.day:02d}）'
    pat = r'[（(]\s*\d{1,2}\s*[.．]\s*\d{1,2}\s*[）)]'
    if re.search(pat, old):
        new = re.sub(pat, stamp, old, count=1)
    else:
        new = old + stamp
    ws.cell(row=1, column=1).value = new
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), new


def ensure_template(cls: dict, src: Path) -> None:
    """邮件没带空白模板时，用技能内置模板兜底，并把标题日期改成业务日。就地改 cls。"""
    if not cls['ok'] or 'template' in cls['roles']:
        return
    biz = (cls['roles'].get('staff_cur', {}).get('data_date')
           or cls['roles'].get('auth', {}).get('data_date')
           or cls['roles'].get('donghu_day', {}).get('data_date'))
    if not biz:
        cls['ok'] = False
        cls['problems'].append('邮件没带模板，且算不出业务日，无法用内置模板兜底')
    elif not FALLBACK_TEMPLATE.exists():
        cls['ok'] = False
        cls['problems'].append(f'邮件没带模板，内置模板也缺失：{FALLBACK_TEMPLATE}')
    else:
        blob, title = _retitle(FALLBACK_TEMPLATE.read_bytes(), biz)
        fname = f'{title}.xlsx'
        _save(src / fname, blob)
        entry = probe_template(blob) or {'kind': 'template', 'labels': [], 'title': title}
        entry.update({'filename': fname, 'blob': blob, 'fallback': True, 'title_date': biz})
        cls['roles']['template'] = entry
        cls['notices'].append(f'使用内置模板兜底，标题日期按业务日 {biz} 写成「{title}」')


def _brief(m: dict) -> dict:
    return {'uid': m['uid'], 'subject': m['subject'], 'from': m['from'],
            'received_date': m.get('received_date')}


def cmd_find(args) -> dict:
    """按约定找今天的源邮件。

    约定：主题含「线上服务」且当天收到 → 直接执行（auto=true）。
    不符合就进一步检索（按附件内容、按最近日期逐级放宽），
    只报告结论并给出建议，**不自动执行**，由用户决定以哪封为准。
    """
    today = date.fromisoformat(args.date) if args.date else datetime.now(TZ).date()
    keyword = args.keyword
    acct = resolve_account()
    imap = connect(acct)
    try:
        typ, _ = imap.select(acct['inbox'], readonly=True)
        if typ != 'OK':
            raise RuntimeError(f"SELECT {acct['inbox']} 失败")
        recent = list_recent(imap, today - timedelta(days=args.days))
        cache: dict = {}
        kw = [m for m in recent if keyword in (m['subject'] or '')]
        kw_today = [m for m in kw if m['received_date'] == today.isoformat()]

        base = {'ok': True, 'keyword': keyword, 'today': today.isoformat(),
                'inbox': acct['inbox'], 'scanned': len(recent),
                'keyword_hits': [_brief(m) for m in kw],
                'keyword_hits_today': [_brief(m) for m in kw_today]}

        # 约定命中：主题带关键词 + 当天收到 + 附件齐 → 直接执行
        if kw_today:
            chk = pool_check(imap, [m['uid'] for m in kw_today], cache)
            if chk['complete']:
                base.update({'auto': True, 'uids': [m['uid'] for m in kw_today],
                             'proposed': [_brief(m) for m in kw_today],
                             'verification': strip_blobs(chk),
                             'reason': '主题含关键词且为当天邮件，附件齐全'})
                return base
            why = '；'.join(chk['problems']) or '必需附件不全'
            base['keyword_today_problems'] = chk['problems']
        else:
            why = (f'今天（{today.isoformat()}）没有主题含「{keyword}」的邮件'
                   if not kw else f'主题含「{keyword}」的邮件都不是今天收到的')

        # 不合约定：放宽检索，用贪心挑出**最少**的几封，只报告不执行
        days = sorted({m['received_date'] for m in recent if m['received_date']}, reverse=True)
        scopes = [
            (f'窗口内主题含「{keyword}」的邮件（{args.days} 天）', [m['uid'] for m in kw]),
            (f'最近一天（{days[0]}）收到的全部邮件' if days else '', 
             [m['uid'] for m in recent if m['received_date'] == days[0]] if days else []),
            ('窗口内收到的全部邮件', [m['uid'] for m in recent]),
        ]
        tried, found = [], None
        for label, uids in scopes:
            if not uids:
                continue
            pick, covered, _ = minimal_pick(imap, uids, cache)
            chk = pool_check(imap, pick, cache) if pick else {'complete': False, 'problems': ['没有可用邮件'],
                                                             'roles': {}, 'notices': []}
            tried.append({'scope': label, 'scanned': len(uids), 'picked': pick,
                          'complete': chk.get('complete', False),
                          'roles': sorted(chk.get('roles', {})),
                          'missing': sorted(set(REQUIRED_ROLES) - covered)})
            if chk.get('complete'):
                # 先严后宽，第一个凑齐的就是最贴近约定的，后面的范围只会更松，不必再下载
                found = {'scope': label, 'uids': pick, 'verification': chk}
                break

        base['auto'] = False
        base['reason'] = why
        base['tried'] = tried
        if found:
            uids = found['uids']
            base['uids'] = uids
            base['proposed'] = [_brief(cache[u]) for u in uids]
            base['verification'] = strip_blobs(found['verification'])
            base['question'] = (
                f'{why}。按附件内容检索（{found["scope"]}），'
                f'其中 {len(uids)} 封可以凑齐全部必需附件。是否按这几封继续？'
                f'（uid：{"、".join(uids)}；以哪封为准也可以直接指定）')
        else:
            missing = sorted({r for t in tried for r in t['missing']})
            base['question'] = (f'{why}，且放宽检索也没能凑齐必需附件。'
                               f'已检索 {len(recent)} 封、{args.days} 天内的邮件。'
                               f'还缺：{"、".join(missing) or "未知"}。'
                               f'请确认源邮件是否已到，或直接指定邮件。')
            base['missing_roles'] = missing
        return base
    finally:
        try:
            imap.logout()
        except Exception:
            pass


def cmd_fetch(args) -> dict:
    acct = resolve_account()
    imap = connect(acct)
    try:
        typ, _ = imap.select(acct['inbox'], readonly=True)
        if typ != 'OK':
            raise RuntimeError(f"SELECT {acct['inbox']} 失败")
        msgs = [fetch_message(imap, u) for u in args.uid]
    finally:
        try:
            imap.logout()
        except Exception:
            pass

    dated = [m for m in msgs if m['received_date']]
    if not dated:
        return {'ok': False, 'error': '取不到邮件日期，无法命名工作目录'}
    # 主邮件 = 带「汇总通报表样式」模板的那封（即「日通报表」邮件）；
    # 找不到就退回最新的一封。回信对象与工作目录日期都以它为准。
    _all = classify([a for m in msgs for a in m['attachments']])
    _tpl = _all['roles'].get('template', {}).get('filename')
    main = next((m for m in dated if any(a['filename'] == _tpl for a in m['attachments'])),
                max(dated, key=lambda m: m['received']))
    workdir = Path(args.workdir) if args.workdir else _workdir_for(Path(args.root), main['received'].date())
    src = workdir / 'source'
    attachments = []
    for m in msgs:
        for att in m['attachments']:
            _save(src / att['filename'], att['blob'])
            attachments.append(att)
    cls = classify(attachments)
    ensure_template(cls, src)
    meta = {
        'ok': cls['ok'],
        'workdir': str(workdir),
        'source_dir': str(src),
        'received_date': main['received_date'],
        'mails': [{'uid': m['uid'], 'subject': m['subject'], 'from': m['from'],
                   'received_date': m['received_date'],
                   'attachments': [a['filename'] for a in m['attachments']]} for m in msgs],
        'roles': {k: {'file': v['filename'], 'data_date': v.get('data_date')}
                  for k, v in cls['roles'].items()},
        'unknown_files': cls['unknown_files'],
        'problems': cls['problems'],
        'notices': cls.get('notices', []),
        'attachment_count': cls.get('attachment_count', len(attachments)),
        'reply': {'to': parseaddr(main['from'])[1], 'subject': main['subject'],
                  'message_id': main['message_id'], 'uid': main['uid']},
    }
    # 源表识别的完整结果落盘，build 直接复用，避免二次解析
    payload = {k: {'filename': v['filename'], 'data_date': v.get('data_date'),
                   'nets': v.get('nets'), 'total': v.get('total'),
                   'title': v.get('title'), 'labels': v.get('labels')}
               for k, v in cls['roles'].items()}
    _save(workdir / 'roles.json', json.dumps({'roles': payload, 'reply': meta['reply']},
                                             ensure_ascii=False, indent=1).encode('utf-8'))
    return meta


def _load_roles(workdir: Path, source: Path | None) -> dict:
    meta_p = workdir / 'roles.json'
    if not meta_p.exists():
        raise RuntimeError(f'缺少 {meta_p}，请先执行 fetch')
    meta = json.loads(meta_p.read_text('utf-8'))
    src = source or (workdir / 'source')
    roles = {}
    for role, info in meta['roles'].items():
        blob = (src / info['filename']).read_bytes()
        entry = dict(info)
        if role == 'template':
            entry['blob'] = blob
        roles[role] = entry
    return {'roles': roles, 'reply': meta.get('reply', {}),
            'report_date': meta.get('report_date')}


def _remember_report_date(workdir: Path, day_date) -> None:
    """把通报日期补进 roles.json，send 据此组回信标题（格式：线上服务_YYYYMMDD）。"""
    meta_p = workdir / 'roles.json'
    if not meta_p.exists():
        return
    try:
        meta = json.loads(meta_p.read_text('utf-8'))
    except Exception:
        return
    meta['report_date'] = day_date
    meta_p.write_text(json.dumps(meta, ensure_ascii=False, indent=1), 'utf-8')


def reply_subject(day_date) -> str:
    """回信标题：线上服务_YYYYMMDD（通报日期）。"""
    digits = re.sub(r'\D', '', str(day_date or ''))
    return f'线上服务_{digits}' if len(digits) == 8 else '线上服务'


def cmd_build(args) -> dict:
    workdir = Path(args.workdir)
    loaded = _load_roles(workdir, Path(args.source) if args.source else None)
    roles = loaded['roles']
    for need in ('donghu_day', 'donghu_prev', 'staff_cur', 'staff_prev', 'staff_month', 'auth', 'template'):
        if need not in roles:
            return {'ok': False, 'error': f'缺少角色 {need}，无法计算；请检查源表或用 fetch 重新取信'}
    data = compute(roles)
    out_dir = workdir / 'out'
    title = roles['template']['title']
    out_path = out_dir / f'{title}.xlsx'
    write_report(roles['template']['blob'], data, out_path)

    payload = {'day_date': data['day_date'], 'prev_date': data['prev_date'],
               'rows': data['rows'], 'warnings': data['warnings']}
    text_info = build_text(data) if not args.no_text else {'text': '', 'notes': []}
    (workdir / 'data.json').write_text(json.dumps(payload, ensure_ascii=False, indent=1), 'utf-8')
    # 把通报日期写回 roles.json，send 要用它组回信标题
    _remember_report_date(workdir, data['day_date'])
    if text_info['text']:
        txt = workdir / f'文案-{data["day_date"][5:].replace("-", "")}.txt'
        txt.write_text(text_info['text'], 'utf-8')
    else:
        txt = None
    return {'ok': True, 'workdir': str(workdir), 'report': str(out_path),
            'report_name': out_path.name, 'title': title,
            'text_file': str(txt) if txt else None,
            'text': text_info['text'], 'notes': text_info['notes'],
            'day_date': data['day_date'], 'prev_date': data['prev_date'],
            'staff_dates': data['staff_dates'], 'auth_date': data['auth_date'],
            'warnings': data['warnings'],
            'reply': loaded['reply']}


def cmd_send(args) -> dict:
    workdir = Path(args.workdir)
    loaded = _load_roles(workdir, None)
    reply = loaded['reply']
    out_dir = workdir / 'out'
    xlsx = sorted(out_dir.glob('*.xlsx'))
    if not xlsx:
        return {'ok': False, 'error': f'{out_dir} 下没有 xlsx，请先 build'}
    report = xlsx[0]
    txts = sorted(workdir.glob('文案-*.txt'))
    if not txts:
        return {'ok': False, 'error': '没有文案文件，请先 build'}
    body = txts[-1].read_text('utf-8')
    sender = (args.to or reply.get('to') or '').strip()
    if not sender:
        return {'ok': False, 'error': '没有收件人（源邮件没有解析出发件人地址）'}
    # 固定抄送名单 + 原发件人；发件人本身在名单里就只发一遍
    to_list = dedup_addresses([sender] + list(FIXED_RECIPIENTS))
    to = ', '.join(to_list)
    # 回信标题统一为「线上服务_YYYYMMDD」（通报日期），不再沿用原邮件主题
    subject = args.subject or reply_subject(loaded.get('report_date'))

    acct = resolve_account()
    msg = EmailMessage()
    msg['From'] = acct['user']
    msg['To'] = to
    msg['Subject'] = subject
    msg['Date'] = formatdate(localtime=True)
    if reply.get('message_id'):
        msg['In-Reply-To'] = reply['message_id']
        msg['References'] = reply['message_id']
    msg.set_content(body)
    msg.add_attachment(report.read_bytes(),
                       maintype='application',
                       subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                       filename=report.name)
    if args.dry_run:
        return {'ok': True, 'dry_run': True, 'to': to, 'recipients': to_list, 'subject': subject,
                'attachment': report.name, 'bytes': len(report.read_bytes()),
                'smtp': f"{acct['smtp_host']}:{acct['smtp_port']}", 'body': body}

    force_ipv4()
    ctx = ssl.create_default_context()
    if acct['smtp_port'] == 465:
        server = smtplib.SMTP_SSL(acct['smtp_host'], acct['smtp_port'], context=ctx, timeout=60)
    else:
        server = smtplib.SMTP(acct['smtp_host'], acct['smtp_port'], timeout=60)
        server.starttls(context=ctx)
    try:
        server.login(acct['user'], acct['password'])
        server.send_message(msg)
    finally:
        try:
            server.quit()
        except Exception:
            pass
    return {'ok': True, 'sent': True, 'to': to, 'recipients': to_list, 'subject': subject,
            'attachment': report.name, 'smtp': f"{acct['smtp_host']}:{acct['smtp_port']}"}


def cmd_run(args) -> dict:
    got = cmd_fetch(args)
    if not got.get('ok'):
        return got
    out = cmd_build(argparse.Namespace(workdir=got['workdir'], source=None, no_text=False))
    # 取信阶段的关键结论提到顶层，免得模型还要往下翻一层
    for k in ('received_date', 'attachment_count', 'problems', 'notices',
              'unknown_files', 'roles', 'workdir', 'mails'):
        out.setdefault(k, got.get(k))
    out['fetch'] = got
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description='网点线上服务能力提升通报：取信/计算/出表/回信')
    sub = ap.add_subparsers(dest='cmd', required=True)

    d = sub.add_parser('find', help='按约定找源邮件：主题含关键词且当天收到则直接执行，否则报告并请示')
    d.add_argument('--keyword', default=KEYWORD, help=f'主题关键词，默认「{KEYWORD}」')
    d.add_argument('--date', default=None, help='把哪天当作"今天"（YYYY-MM-DD），默认系统当天')
    d.add_argument('--days', type=int, default=3, help='放宽检索的回看天数，默认 3')
    d.set_defaults(func=cmd_find)

    f = sub.add_parser('fetch', help='按 uid 取信并把附件存到工作目录')
    f.add_argument('--uid', action='append', required=True, help='源邮件 uid，可重复')
    f.add_argument('--root', default=str(DEFAULT_ROOT), help='工作目录根，默认 D:\\资料\\ICBC\\automations')
    f.add_argument('--workdir', default=None, help='直接指定工作目录（默认按收件日新建）')
    f.set_defaults(func=cmd_fetch)

    b = sub.add_parser('build', help='计算并写出通报表与文案')
    b.add_argument('--workdir', required=True)
    b.add_argument('--source', default=None, help='源表目录，默认 <workdir>/source')
    b.add_argument('--no-text', action='store_true')
    b.set_defaults(func=cmd_build)

    s = sub.add_parser('send', help='把通报表与文案回复给源邮件发件人')
    s.add_argument('--workdir', required=True)
    s.add_argument('--to', default=None, help='覆盖收件人')
    s.add_argument('--subject', default=None, help='覆盖回信标题，默认「线上服务_YYYYMMDD」')
    s.add_argument('--dry-run', action='store_true')
    s.set_defaults(func=cmd_send)

    r = sub.add_parser('run', help='fetch + build（不发信）')
    r.add_argument('--uid', action='append', required=True)
    r.add_argument('--root', default=str(DEFAULT_ROOT))
    r.add_argument('--workdir', default=None)
    r.set_defaults(func=cmd_run)

    args = ap.parse_args()
    try:
        out = args.func(args)
    except Exception as exc:
        import traceback
        out = {'ok': False, 'error': f'{type(exc).__name__}: {exc}',
               'traceback': traceback.format_exc()[-1500:]}
    print(json.dumps(out, ensure_ascii=False, indent=1, default=str))
    raise SystemExit(0 if out.get('ok') else 1)


if __name__ == '__main__':
    main()
