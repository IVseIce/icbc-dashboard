#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""fivecards.py —— 五卡汇总表（支行维度）解析。

源表：第 1 行卡片名（合并单元格，需向右填充），第 2 行指标名，3–10 行支行，末行分行合计。
输出：data/snapshots/wk/<YYYY-MM-DD>.json
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl

PROJECT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(PROJECT / 'src'))

import viz  # noqa: E402


def num(v):
    if v is None or v == '':
        return None
    try:
        f = float(str(v).replace(',', ''))
        return f
    except ValueError:
        return None


def parse(path: Path, as_of: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            for r in range(1, ws.max_row + 1)]
    card_row, metric_row = rows[0], rows[1]

    cards: list[str] = []
    cur = ''
    for v in card_row[1:]:
        if v not in (None, ''):
            cur = str(v).strip()
        cards.append(cur)

    cols = []  # (card, metric, col_index)
    for i, (card, metric) in enumerate(zip(cards, metric_row[1:]), start=1):
        if not metric or str(metric).strip() == '':
            continue
        cols.append((card, str(metric).strip(), i))

    branches = []
    total = None
    for r in rows[2:]:
        name = str(r[0]).strip() if r[0] else ''
        if not name:
            continue
        vals = {}
        for card, metric, i in cols:
            v = num(r[i])
            if v is not None:
                vals[f'{card}|{metric}'] = v
        rec = {'name': name, 'values': vals}
        if '分行' in name or name in ('合计', '总计'):
            total = rec
        else:
            branches.append(rec)

    wb.close()
    return {'as_of': as_of, 'cards': [c for c in dict.fromkeys(cards) if c],
            'metrics': sorted({m for _, m, _ in cols}),
            'cols': [{'card': c, 'metric': m} for c, m, _ in cols],
            'branches': branches, 'total': total}


def main() -> None:
    src = Path(sys.argv[1] if len(sys.argv) > 1 else '/mnt/d/资料/ICBC/20260909五卡/五卡汇总表.xlsx')
    as_of = sys.argv[2] if len(sys.argv) > 2 else '2026-09-09'
    snap = parse(src, as_of)
    f = viz.paths()['snapshots'] / 'wk' / f'{as_of}.json'
    f.parent.mkdir(parents=True, exist_ok=True)
    f.write_text(json.dumps(snap, ensure_ascii=False), encoding='utf-8')
    print(json.dumps({'ok': True, 'file': str(f), 'cards': snap['cards'],
                      'branches': len(snap['branches'])}, ensure_ascii=False))


if __name__ == '__main__':
    main()
